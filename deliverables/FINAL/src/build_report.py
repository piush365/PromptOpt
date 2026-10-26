"""Build MP1_Report_PromptOpt.docx on top of the college template (front matter kept), two passes for the TOC pages."""
import copy, json, os, re, subprocess, sys
from PIL import Image
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import report_content as C

HERE = os.path.dirname(os.path.abspath(__file__))
TEMPLATE = "/mnt/SharedDrive/Projects/PromptOpt/Mini Project 1_report format 26-27.docx"
FIG = os.path.join(HERE, "..", "figures")
W = qn

# ------------------------------------------------------------------ low-level helpers
def clear_runs(p):
    for ch in list(p):
        if ch.tag in (W("w:r"), W("w:hyperlink"), W("w:proofErr"), W("w:bookmarkStart"), W("w:bookmarkEnd")):
            p.remove(ch)

def first_rpr(p):
    r = p.find(".//" + W("w:r"))
    rpr = r.find(W("w:rPr")) if r is not None else None
    return copy.deepcopy(rpr) if rpr is not None else OxmlElement("w:rPr")

def set_runs(p, segs, rpr=None):
    """segs: [(text, bold)]; bold None keeps the template run's weight."""
    base = rpr if rpr is not None else first_rpr(p)
    clear_runs(p)
    for text, bold in segs:
        r = OxmlElement("w:r"); rp = copy.deepcopy(base)
        if bold is not None:
            for b in rp.findall(W("w:b")) + rp.findall(W("w:bCs")):
                rp.remove(b)
            if bold:
                b = OxmlElement("w:b"); rp.insert(1 if len(rp) else 0, b)
        r.append(rp)
        t = OxmlElement("w:t"); t.text = text; t.set(qn("xml:space"), "preserve"); r.append(t)
        p.append(r)

def set_text(p, text, bold=None):
    set_runs(p, [(text, bold)])

def ptext(el):
    return "".join(t.text or "" for t in el.iter(W("w:t")))

def font(run, size=11, bold=False, italic=False, name="Times New Roman", color=None):
    run.font.name = name; run.font.size = Pt(size); run.font.bold = bold; run.font.italic = italic
    rpr = run._r.get_or_add_rPr(); rf = rpr.find(W("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.insert(0, rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), name)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)

def shade(el_pr, fill):
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill)
    el_pr.append(shd)

def outline(p, level):
    ppr = p._p.get_or_add_pPr(); o = OxmlElement("w:outlineLvl"); o.set(qn("w:val"), str(level)); ppr.append(o)

# ------------------------------------------------------------------ build
def build(pages, out_path):
    d = Document(TEMPLATE)
    body = d.element.body
    els = list(body.iterchildren())
    assert ptext(els[18]) == "<Project Name>" and ptext(els[103]) == "Abstract", "template layout changed"

    # cover
    set_text(els[18], C.TITLE)
    names = {"<Name1>": C.TEAM[0][0], "<PRN1>": C.TEAM[0][1], "<Name2>": C.TEAM[1][0], "<PRN2>": C.TEAM[1][1],
             "<Name3>": C.TEAM[2][0], "<PRN3>": C.TEAM[2][1]}
    for p in els[23].iter(W("w:p")):
        t = ptext(p).strip()
        if t in names:
            set_text(p, names[t])
    set_text(els[28], C.GUIDE)

    # certificate
    students = f"{C.TEAM[0][0]}, {C.TEAM[1][0]} and {C.TEAM[2][0]}"
    set_runs(els[53], [("This is to certify that the Project Report entitled, ", False), ("“" + C.TITLE_UP + "”", True),
                       (" submitted by ", False), (students, True),
                       (", to Walchand College of Engineering Sangli, India, is a record of bonafide project work of course ", False),
                       ("Mini Project-I (7CS345)", True),
                       (" carried out by them under my/our supervision and guidance and is worthy of consideration for the "
                        "award of the degree of Bachelor of Technology in Computer Science & Engineering of the Institute.", False)])
    for p in els[61].iter(W("w:p")):
        if "<Name of guide>" in ptext(p):
            set_text(p, C.GUIDE)

    # acknowledgement: three paragraphs in the template's first paragraph format
    anchor = els[65]; set_text(anchor, C.ACK[0])
    for t in C.ACK[1:]:
        n = copy.deepcopy(els[65]); set_text(n, t); anchor.addnext(n); anchor = n
    body.remove(els[66])

    # declaration
    set_runs(els[70], [("We hereby declare that the work presented in this project report titled ", False),
                       ("“" + C.TITLE_UP + "”", True), (" submitted by us in the partial fulfillment of the requirement of the award of the degree of ", False),
                       ("Bachelor of Technology (B. Tech)", True), (" in the ", False),
                       ("Department of Computer Science & Engineering, Walchand College of Engineering, Sangli", True),
                       (" is an authentic record of our project work carried out under the guidance of " + C.GUIDE + ".", False)])
    set_text(els[85], f"{C.TEAM[0][0]} ({C.TEAM[0][1]})")
    anchor = els[85]
    for nm, prn in C.TEAM[1:]:
        n = copy.deepcopy(els[85]); set_text(n, f"{nm} ({prn})"); anchor.addnext(n); anchor = n
    set_text(els[86], "")

    # keep copies, drop the template's TOC / list / body
    heading_tpl = copy.deepcopy(els[117])
    abstract_h = copy.deepcopy(els[103])
    for i in list(range(88, 102)) + list(range(103, 314)):
        if i in (90, 97):
            continue
        body.remove(els[i])
    # 102: last acronym line carries the section break -> empty it, no bullet
    ppr = els[102].find(W("w:pPr")); num = ppr.find(W("w:numPr"))
    if num is not None:
        ppr.remove(num)
    clear_runs(els[102])
    lof_h, acr_h = els[90], els[97]
    for i in (19, 24, 32, 37, 38, 39, 55, 56):  # blank lines: keep cover and certificate on one page each
        body.remove(els[i])
    pb = OxmlElement("w:pageBreakBefore"); lof_h.find(W("w:pPr")).insert(1, pb)

    sect_last = body[-1]
    def add_el(el):
        sect_last.addprevious(el)

    def new_par():
        return d.add_paragraph()  # appended before the final sectPr

    toc_entries = []   # (level, key, label)
    figs, tbls = [], []  # (label, caption)
    counters = {"fig": 0, "tbl": 0, "ch": 0}
    BODY_PT, LINE = 12, 1.5

    def body_p(text, size=BODY_PT, align=WD_ALIGN_PARAGRAPH.JUSTIFY, bold=False, italic=False, after=6, before=0, line=LINE):
        p = new_par(); p.alignment = align
        pf = p.paragraph_format; pf.space_after = Pt(after); pf.space_before = Pt(before); pf.line_spacing = line
        pf.widow_control = True
        font(p.add_run(text), size, bold, italic); return p

    def border(p_el, side="bottom", color="808080", sz="8", space="4"):
        ppr = p_el.get_or_add_pPr() if hasattr(p_el, "get_or_add_pPr") else p_el.find(W("w:pPr"))
        bdr = OxmlElement("w:pBdr"); b = OxmlElement(f"w:{side}")
        for k, v in (("w:val", "single"), ("w:sz", sz), ("w:space", space), ("w:color", color)):
            b.set(qn(k), v)
        bdr.append(b)
        # pBdr must come before spacing/ind/jc in pPr
        anchor_tags = [W(t) for t in ("w:shd", "w:tabs", "w:spacing", "w:ind", "w:jc", "w:outlineLvl", "w:rPr")]
        for i, ch in enumerate(list(ppr)):
            if ch.tag in anchor_tags:
                ppr.insert(i, bdr); break
        else:
            ppr.append(bdr)

    def chapter(title):
        counters["ch"] += 1; counters["fig"] = 0; counters["tbl"] = 0
        p = copy.deepcopy(heading_tpl)
        for b in p.findall(W("w:bookmarkStart")) + p.findall(W("w:bookmarkEnd")):
            p.remove(b)
        rpr = first_rpr(p)
        for tag in ("w:sz", "w:szCs"):
            e = rpr.find(W(tag))
            if e is not None:
                e.set(qn("w:val"), "32")
        b = OxmlElement("w:b"); rpr.insert(1, b)
        set_runs(p, [(title, None)], rpr)
        ppr = p.find(W("w:pPr")); ppr.insert(1, OxmlElement("w:pageBreakBefore"))
        sp = ppr.find(W("w:spacing"))
        if sp is None:
            sp = OxmlElement("w:spacing"); ppr.append(sp)
        sp.set(qn("w:after"), "280"); sp.set(qn("w:before"), "0")
        border(p, color="1F3864", sz="12", space="6")
        add_el(p)
        toc_entries.append((1, f"ch{counters['ch']}", f"{counters['ch']}  {title}"))

    def h2(text, toc=True):
        p = body_p(text, 13, WD_ALIGN_PARAGRAPH.LEFT, bold=True, after=6, before=14, line=1.15)
        p.paragraph_format.keep_with_next = True; outline(p, 1)
        p.runs[0].font.color.rgb = RGBColor.from_string("1F3864")
        if toc:
            toc_entries.append((2, "h2:" + text, text))

    def bullets(items, numbered=False):
        for i, t in enumerate(items, 1):
            p = new_par(); pf = p.paragraph_format
            pf.left_indent = Inches(0.4); pf.first_line_indent = Inches(-0.25); pf.space_after = Pt(4); pf.line_spacing = 1.4
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            pf.tab_stops.add_tab_stop(Inches(0.4))
            font(p.add_run((f"{i}." if numbered else "•") + "\t"), BODY_PT)
            font(p.add_run(t), BODY_PT)

    def caption_par(label, text, kind, container=None):
        p = container if container is not None else new_par()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf = p.paragraph_format; pf.line_spacing = 1.0
        pf.space_before = Pt(4 if kind == "fig" else 10); pf.space_after = Pt(12 if kind == "fig" else 4)
        font(p.add_run(f"{label}: "), 10.5, bold=True); font(p.add_run(text), 10.5, italic=(kind == "fig"))
        if kind == "tbl":
            pf.keep_with_next = True
        return p

    def next_label(kind):
        counters[kind] += 1
        return ("Figure" if kind == "fig" else "Table") + f" {counters['ch']}.{counters[kind]}"

    def pic_border(run):
        for sppr in run._r.iter(qn("pic:spPr")):
            ln = OxmlElement("a:ln"); ln.set("w", "9525"); sf = OxmlElement("a:solidFill"); c = OxmlElement("a:srgbClr")
            c.set("val", "A6A6A6"); sf.append(c); ln.append(sf); sppr.append(ln)

    def fit(fname, width, max_h):
        w, h = Image.open(os.path.join(FIG, fname)).size
        width = min(width, 6.5)
        if width * h / w > max_h:
            width = max_h * w / h
        return width

    def figure(fname, cap, width, framed=None):
        framed = fname.startswith(("ui_", "gh_")) if framed is None else framed
        width = fit(fname, width, 7.0)
        p = new_par(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_before = Pt(8); p.paragraph_format.space_after = Pt(0)
        r = p.add_run(); r.add_picture(os.path.join(FIG, fname), width=Inches(width))
        if framed:
            pic_border(r)
        label = next_label("fig"); caption_par(label, cap, "fig"); figs.append((label, cap))

    def figpair(f1, c1, f2, c2, max_h=3.6):
        t = d.add_table(rows=1, cols=2); t.alignment = WD_TABLE_ALIGNMENT.CENTER
        tblPr = t._tbl.tblPr; b = OxmlElement("w:tblBorders")
        for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
            e = OxmlElement(f"w:{side}"); e.set(qn("w:val"), "nil"); b.append(e)
        tblPr.append(b)
        trpr = t.rows[0]._tr.get_or_add_trPr(); cs = OxmlElement("w:cantSplit"); trpr.append(cs)
        for c_i, (fname, cap) in enumerate(((f1, c1), (f2, c2))):
            cell = t.rows[0].cells[c_i]; cell.width = Inches(3.25)
            p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(6)
            r = p.add_run(); r.add_picture(os.path.join(FIG, fname), width=Inches(fit(fname, 3.05, max_h))); pic_border(r)
            label = next_label("fig"); caption_par(label, cap, "fig", cell.add_paragraph()); figs.append((label, cap))
        for gc in t._tbl.tblGrid.findall(W("w:gridCol")):
            gc.set(qn("w:w"), str(int(3.25 * 1440)))
        body_p("", 4, after=0, line=1.0)

    def table(cap, header, rows, widths, size=10, status_col=None):
        label = next_label("tbl"); caption_par(label, cap, "tbl"); tbls.append((label, cap))
        t = d.add_table(rows=1 + len(rows), cols=len(header)); t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        tblPr = t._tbl.tblPr; bd = OxmlElement("w:tblBorders")
        for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
            e = OxmlElement(f"w:{side}"); e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4"); e.set(qn("w:color"), "A6A6A6")
            bd.append(e)
        tblPr.append(bd)
        mar = OxmlElement("w:tblCellMar")
        for side, v in (("top", 40), ("bottom", 40), ("left", 90), ("right", 90)):
            e = OxmlElement(f"w:{side}"); e.set(qn("w:w"), str(v)); e.set(qn("w:type"), "dxa"); mar.append(e)
        tblPr.append(mar)
        tot = sum(widths); widths = [w * 6.5 / tot for w in widths] if tot > 6.5 else widths
        for r_i, row in enumerate([header] + rows):
            tr = t.rows[r_i]; trpr = tr._tr.get_or_add_trPr()
            trpr.append(OxmlElement("w:cantSplit"))
            if r_i == 0:
                trpr.append(OxmlElement("w:tblHeader"))
            for c_i, val in enumerate(row):
                cell = tr.cells[c_i]; cell.width = Inches(widths[c_i])
                p = cell.paragraphs[0]; pf = p.paragraph_format
                pf.space_after = Pt(0); pf.space_before = Pt(0); pf.line_spacing = 1.0
                color = None
                if status_col is not None and c_i == status_col and r_i > 0:
                    color = "1E7B34" if val == "PASS" else "B42318"
                font(p.add_run(str(val)), size, bold=(r_i == 0 or color is not None), color="FFFFFF" if r_i == 0 else color)
                if r_i == 0:
                    shade(cell._tc.get_or_add_tcPr(), "1F3864")
                elif r_i % 2 == 0:
                    shade(cell._tc.get_or_add_tcPr(), "F2F5FA")
        grid = t._tbl.tblGrid
        for gc, w in zip(grid.findall(W("w:gridCol")), widths):
            gc.set(qn("w:w"), str(int(w * 1440)))
        body_p("", 6, after=4, line=1.0)

    def h3(text):
        p = body_p(text, 12, WD_ALIGN_PARAGRAPH.LEFT, bold=True, after=4, before=8, line=1.15)
        p.paragraph_format.keep_with_next = True
        p.runs[0].font.color.rgb = RGBColor.from_string("2F4B7C")

    def eq(text):
        p = body_p("", after=8, before=4, align=WD_ALIGN_PARAGRAPH.CENTER, line=1.2)
        font(p.add_run(text), 11.5, italic=True)

    def code(text):
        lines = text.split("\n")
        for i, ln in enumerate(lines):
            p = new_par(); pf = p.paragraph_format; pf.space_after = Pt(0); pf.space_before = Pt(0); pf.line_spacing = 1.0
            pf.left_indent = Inches(0.15); pf.right_indent = Inches(0.15); pf.keep_with_next = i < len(lines) - 1
            shade(p._p.get_or_add_pPr(), "F2F2F2")
            font(p.add_run(ln if ln else " "), 9, name="Courier New")
        body_p("", 6, after=4, line=1.0)

    # ---------------- Abstract
    add_el(abstract_h)
    toc_entries.append((0, "abstract", "Abstract"))
    for t in C.ABSTRACT:
        body_p(t, after=10)
    kp = body_p("", after=6)
    font(kp.add_run("Keywords: "), BODY_PT, bold=True)
    font(kp.add_run(C.KEYWORDS), BODY_PT, italic=True)

    # ---------------- chapters
    for title, blocks in C.CH:
        chapter(title)
        for b in blocks:
            k = b[0]
            if k == "p":
                pp = body_p(b[1]); pp.paragraph_format.keep_with_next = b[1].endswith(":")
            elif k == "h2": h2(b[1])
            elif k == "h3": h3(b[1])
            elif k == "eq": eq(b[1])
            elif k == "bul": bullets(b[1])
            elif k == "num": bullets(b[1], numbered=True)
            elif k == "tbl": table(b[1], b[2], b[3], b[4])
            elif k == "tc": table(b[1], b[2], b[3], b[4], size=8.5, status_col=6)
            elif k == "fig": figure(b[1], b[2], b[3])
            elif k == "pair": figpair(*b[1:])
            elif k == "code": code(b[1])
            else: raise ValueError(k)

    chapter("References")
    for i, r in enumerate(C.REFS, 1):
        p = body_p(f"[{i}]\t{r}", 11.5, after=4, line=1.15, align=WD_ALIGN_PARAGRAPH.LEFT)
        p.paragraph_format.left_indent = Inches(0.45); p.paragraph_format.first_line_indent = Inches(-0.45)
        p.paragraph_format.tab_stops.add_tab_stop(Inches(0.45))
    chapter("Scholarly contribution")
    for head, ans in C.SCHOLARLY:
        h2(head, toc=False); body_p(ans)

    # ---------------- running header / footer for the body (Abstract onwards) only
    sec = d.sections[-1]
    sec.header.is_linked_to_previous = False; sec.footer.is_linked_to_previous = False
    hp = sec.header.paragraphs[0]; hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    font(hp.add_run("PromptOpt: Token-Efficient Prompt Optimization for LLMs"), 9, italic=True, color="595959")
    border(hp._p, color="A6A6A6", sz="4", space="1")
    fp = sec.footer.paragraphs[0]
    border(fp._p, side="top", color="A6A6A6", sz="4", space="4")
    fp.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)
    font(fp.add_run("Department of Computer Science and Engineering, WCE Sangli"), 9, color="595959")
    font(fp.add_run("\t"), 9)
    r = fp.add_run(); font(r, 9, color="595959")
    for tag, txt in (("begin", None), (None, " PAGE "), ("end", None)):
        if tag:
            fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), tag); r._r.append(fc)
        else:
            it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = txt; r._r.append(it)

    # ---------------- TOC, list of figures / tables, acronyms (static, page numbers from the previous pass)
    toc_h = els[87]; anchor = toc_h
    def toc_line(text, page, level, after_el, size=11, bold=False):
        p = new_par(); p._p.getparent().remove(p._p)
        pf = p.paragraph_format; pf.space_after = Pt(3 if level < 2 else 1); pf.space_before = Pt(6 if level < 2 else 0)
        pf.left_indent = Inches(0.3 * max(level - 1, 0)); pf.line_spacing = 1.0
        pf.tab_stops.add_tab_stop(Inches(6.45), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        font(p.add_run(text), size, bold=bold); font(p.add_run("\t" + str(page)), size, bold=bold)
        after_el.addnext(p._p); return p._p
    for level, key, label in toc_entries:
        anchor = toc_line(label, pages.get(key, ""), level, anchor, 11.5 if level < 2 else 11, bold=level < 2)
    anchor = lof_h
    for label, cap in figs:
        anchor = toc_line(f"{label}: {cap}", pages.get(label, ""), 2, anchor, 10.5)
    lot_h = copy.deepcopy(lof_h); clear_runs(lot_h)
    r_ = copy.deepcopy(lof_h.find(W("w:r"))); [r_.remove(t) for t in r_.findall(W("w:t"))]
    t_ = OxmlElement("w:t"); t_.text = "List of Tables"; r_.append(t_); lot_h.append(r_)
    anchor.addnext(lot_h); anchor = lot_h
    for label, cap in tbls:
        anchor = toc_line(f"{label}: {cap}", pages.get(label, ""), 2, anchor, 10.5)
    pb2 = OxmlElement("w:pageBreakBefore"); acr_h.find(W("w:pPr")).insert(1, pb2)
    anchor = acr_h
    for ab, full in C.ACRONYMS:
        p = new_par(); p._p.getparent().remove(p._p)
        p.paragraph_format.space_after = Pt(4); p.paragraph_format.tab_stops.add_tab_stop(Inches(1.2))
        font(p.add_run(ab), 11.5, bold=True); font(p.add_run("\t" + full), 11.5)
        anchor.addnext(p._p); anchor = p._p

    d.save(out_path)
    return toc_entries, figs + tbls


def pdf_pages(docx_path):
    outdir = os.path.dirname(docx_path)
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", outdir, docx_path], check=True,
                   stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    pdf = docx_path[:-5] + ".pdf"
    n = int(re.search(r"Pages:\s+(\d+)", subprocess.run(["pdfinfo", pdf], capture_output=True, text=True).stdout).group(1))
    texts = [subprocess.run(["pdftotext", "-layout", "-f", str(i), "-l", str(i), pdf, "-"], capture_output=True,
                            text=True).stdout for i in range(1, n + 1)]
    return pdf, texts


def locate(texts, toc_entries, labelled):
    norm = lambda s: re.sub(r"\s+", " ", s).strip()
    start = next(i for i, t in enumerate(texts) if "Application Programming Interface" in t) + 1
    pages, cur = {}, start
    for level, key, label in toc_entries:
        want = norm(label.split("  ", 1)[1] if level == 1 else label)
        for i in range(cur, len(texts)):
            lines = [norm(l) for l in texts[i].splitlines()]
            if any(l == want or (l.endswith(want) and len(l) <= len(want) + 6) for l in lines):
                pages[key] = i + 1; cur = i; break
        else:
            print("NOT FOUND", label)
    for label, cap in labelled:
        for i in range(start, len(texts)):
            if re.search(re.escape(label) + r":", texts[i]):
                pages[label] = i + 1; break
        else:
            print("NOT FOUND", label)
    return pages


if __name__ == "__main__":
    out = sys.argv[1]
    entries, figs = build({}, out)
    _, texts = pdf_pages(out)
    pages = locate(texts, entries, figs)
    entries, figs = build(pages, out)
    pdf, texts = pdf_pages(out)
    pages2 = locate(texts, entries, figs)
    print("pages stable:", pages == pages2, "| total pages:", len(texts))
    if pages != pages2:
        build(pages2, out); pdf_pages(out)
    json.dump(pages2, open(out + ".pages.json", "w"), indent=1)
