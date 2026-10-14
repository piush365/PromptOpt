"""Fill the LA-2 deck built from the LA-1 Canva slides (same backgrounds, logo, fonts, colours).
Numbers: evaluation/FINAL_RESULTS.md, evaluation/correctness_suite/RESULTS.md, git, test runs on 14 Oct 2026."""
import copy, os, sys
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.oxml.ns import qn
from PIL import Image

HERE = os.path.dirname(os.path.abspath(__file__))
FIG = os.path.join(HERE, "fig")
SRC, OUT = sys.argv[1], sys.argv[2]

MINT, GREEN, DARK, LIGHT, PALE, MAROON, WHITE = "F5FFF9", "457C5B", "1E4926", "9FC4AE", "8AB79C", "8E1711", "FFFFFF"
PASS_C = "1E7B34"
HEAD, BOLD, ULTRA, BODY, MONO = "Manifold Bold", "Manifold Bold", "Manifold Ultra-Bold", "Lato", "Courier New"
rgb = RGBColor.from_string

prs = Presentation(SRC)
S = list(prs.slides)

KEEP = {  # base slide -> shapes to keep (decoration, logo, title)
    "s9": {"Group 5", "Group 8", "Freeform 11", "TextBox 50"},
    "s12": {"Group 2", "Group 5", "Freeform 8", "TextBox 9"},
    "s3": {"Group 2", "Group 5", "Freeform 8", "Freeform 9", "Freeform 44", "TextBox 10"},
    "s16": {"Group 72", "Group 75", "Freeform 78", "Freeform 79", "Freeform 80", "TextBox 81"},
    "s17": {"Group 2", "Group 5", "TextBox 8", "Freeform 9", "Freeform 10", "Freeform 11"},
    "s5": {"Group 2", "Group 5", "TextBox 8", "Freeform 21", "Freeform 22", "Freeform 23"},
}
TITLE = {"s9": "TextBox 50", "s12": "TextBox 9", "s3": "TextBox 10", "s16": "TextBox 81", "s17": "TextBox 8", "s5": "TextBox 8"}
BASES = ["s1", "s2", "s9", "s12", "s12", "s3", "s12", "s9", "s16", "s3", "s12", "s17", "s9", "s12", "s5", "s3", "s16", "s20"]


def set_text_keep(shape, text):
    """Replace a text box's text, keeping the first run's formatting."""
    tf = shape.text_frame; p0 = tf.paragraphs[0]
    r0 = p0.runs[0]
    for r in p0.runs[1:]:
        r._r.getparent().remove(r._r)
    r0.text = text
    for p in tf.paragraphs[1:]:
        p._p.getparent().remove(p._p)


def prep(i, title):
    s = S[i]; base = BASES[i]
    for sh in list(s.shapes):
        if sh.name not in KEEP[base]:
            sh._element.getparent().remove(sh._element)
    t = next(sh for sh in s.shapes if sh.name == TITLE[base])
    set_text_keep(t, title)
    t.text_frame.word_wrap = True
    t.left, t.width = Inches(2.3), Inches(15.4)
    t.top = Inches(0.62)
    return s


def box(s, x, y, w, h, fill, line=None, radius=0.07, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    sh = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = radius
    sh.fill.solid(); sh.fill.fore_color.rgb = rgb(fill)
    if line:
        sh.line.color.rgb = rgb(line); sh.line.width = Pt(1.5)
    else:
        sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def text(s, x, y, w, h, paras, size=20, color=DARK, font=BODY, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, space_after=6, line=1.1):
    """paras: str | list of str | list of list of (text, {opts}) runs."""
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = tb.text_frame
    tf.word_wrap = True; tf.vertical_anchor = anchor
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, 0)
    if isinstance(paras, str):
        paras = [paras]
    for k, para in enumerate(paras):
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        p.alignment = align; p.space_after = Pt(space_after); p.line_spacing = line
        runs = para if isinstance(para, list) else [(para, {})]
        for t, o in runs:
            r = p.add_run(); r.text = t
            f = r.font; f.size = Pt(o.get("size", size)); f.name = o.get("font", font)
            f.bold = o.get("bold", bold); f.color.rgb = rgb(o.get("color", color))
    return tb


def bullets(s, x, y, w, h, items, size=20, color=DARK, dot=GREEN, space_after=8):
    return text(s, x, y, w, h, [[("●  ", {"color": dot, "size": size - 6}), (t, {})] for t in items],
                size=size, color=color, space_after=space_after)


def circle_num(s, x, y, d, n, fill=LIGHT, color=DARK, size=22):
    c = box(s, x, y, d, d, fill, shape=MSO_SHAPE.OVAL)
    text(s, x, y, d, d, str(n), size=size, color=color, font=BOLD, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    return c


def stat(s, x, y, w, h, big, label, fill=GREEN, color=MINT, size=50, lsize=18):
    box(s, x, y, w, h, fill)
    text(s, x + 0.2, y + 0.18, w - 0.4, h * 0.55, big, size=size, color=color, font=ULTRA, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + 0.25, y + h * 0.6, w - 0.5, h * 0.38, label, size=lsize, color=color, align=PP_ALIGN.CENTER)


def picture(s, fname, x, y, w=None, h=None, crop_h=None, border=GREEN):
    path = os.path.join(FIG, fname); iw, ih = Image.open(path).size
    if w is None:
        w = h * iw / ih
    full_h = w * ih / iw
    if crop_h and crop_h < full_h:
        pic = s.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(full_h))
        pic.crop_bottom = 1 - crop_h / full_h; pic.height = Inches(crop_h)
    else:
        pic = s.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(full_h))
    if border:
        pic.line.color.rgb = rgb(border); pic.line.width = Pt(1.5)
    return pic


def style_chart(ch, colors, size=16, legend=True):
    ch.font.name = BODY; ch.font.size = Pt(size); ch.font.color.rgb = rgb(DARK)
    for ser, c in zip(ch.plots[0].series, colors):
        ser.format.fill.solid(); ser.format.fill.fore_color.rgb = rgb(c)
    pl = ch.plots[0]; pl.has_data_labels = True
    dl = pl.data_labels; dl.font.size = Pt(size - 2); dl.font.name = BODY; dl.font.color.rgb = rgb(DARK)
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    dl.number_format = '0'; dl.number_format_is_linked = False
    va = ch.value_axis; va.visible = False; va.has_major_gridlines = False
    ca = ch.category_axis; ca.tick_labels.font.size = Pt(size); ca.format.line.color.rgb = rgb(PALE)
    ch.has_legend = legend
    if legend:
        ch.legend.position = XL_LEGEND_POSITION.BOTTOM; ch.legend.include_in_layout = False; ch.legend.font.size = Pt(size)


def table(s, x, y, w, col_w, rows, size=16, row_h=0.62, status_col=None):
    nr, nc = len(rows), len(rows[0])
    gt = s.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w), Inches(row_h * nr)).table
    tblPr = gt._tbl.tblPr
    for a in ("bandRow", "firstRow"):
        tblPr.set(a, "0")
    sid = tblPr.find(qn("a:tableStyleId"))
    if sid is not None:
        tblPr.remove(sid)
    for j, cw in enumerate(col_w):
        gt.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        gt.rows[i].height = Inches(row_h)
        for j, val in enumerate(row):
            c = gt.cell(i, j); c.fill.solid(); c.fill.fore_color.rgb = rgb(GREEN if i == 0 else (WHITE if i % 2 else MINT))
            c.margin_left = c.margin_right = Inches(0.12); c.margin_top = c.margin_bottom = Inches(0.05)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = c.text_frame; tf.word_wrap = True; p = tf.paragraphs[0]; p.text = ""
            r = p.add_run(); r.text = str(val); f = r.font; f.size = Pt(size)
            f.name = BOLD if i == 0 else BODY; f.bold = i == 0
            color = MINT if i == 0 else DARK
            if status_col is not None and j == status_col and i > 0:
                color = PASS_C if val == "PASS" else MAROON; f.bold = True; f.name = BOLD
            f.color.rgb = rgb(color)
    return gt


def notes(s, t):
    s.notes_slide.notes_text_frame.text = t


P = "Presenter: "

# 1 Title: LA-1's title slide, only LA-1 -> LA-2 (and the 'Sciecne' typo) ----------------------------------
s = S[0]
for sh in s.shapes:
    if sh.has_text_frame:
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                if "LA-1" in r.text:
                    r.text = r.text.replace("LA-1", "LA-2")
                if "Sciecne" in r.text:
                    r.text = r.text.replace("Sciecne", "Science")
    widths = {"TextBox 18": 13.5, "TextBox 19": 11.5, "TextBox 20": 11.0, "TextBox 21": 3.0, "TextBox 11": 9.0, "TextBox 12": 6.0}
    if sh.name in widths:  # widen, same centre, so each line stays on one line in PowerPoint and in Canva
        c = sh.left + sh.width // 2; sh.width = Inches(widths[sh.name]); sh.left = c - sh.width // 2
notes(s, P + "Piush. Greet the panel, introduce the team and our guide Prof. A. S. Pawar. One line: PromptOpt rewrites a "
      "vague prompt into a clear, model-specific one, shows every change, and we measured that it saves tokens. This is Release 2 "
      "for LA-2: what we built since LA-1, how we tested it, and a live demo.")

# 2 Agenda: LA-1's agenda slide with the eight LA-2 topics --------------------------------------------------------
s = S[1]
agenda = {"TextBox 31": "Core Modules", "TextBox 36": "Transaction Flow & Database", "TextBox 33": "Unit Testing",
          "TextBox 37": "Integration & System Testing", "TextBox 34": "Test Cases & Correctness",
          "TextBox 38": "Results & Working Prototype", "TextBox 35": "Git & GitHub", "TextBox 39": "Challenges & Future Plans"}
for sh in s.shapes:
    if sh.name in agenda:
        set_text_keep(sh, agenda[sh.name])
notes(s, P + "Piush. The agenda follows the LA-2 checklist: core modules, transaction flow and database, unit testing, "
      "integration and system testing, test cases, results and the working prototype, Git and GitHub, challenges and plans, "
      "then a live demo.")

# 3 Architecture --------------------------------------------------------------------------------------------------
s = prep(2, "SYSTEM ARCHITECTURE")
box(s, 1.0, 2.3, 18.0, 4.3, WHITE, line=PALE)
picture(s, "01_system_architecture.png", 1.25, 2.5, w=17.5, border=None)
cards = [("Stage A", "spaCy + regex + sentence-embedding classifier: category, missing format and constraints, filler, ambiguous references"),
         ("Stage B", "15 deterministic rules B01-B15, every change logged; category rules only at confidence ≥ 0.6"),
         ("Stage C", "Qwen2.5-0.5B + LoRA, called only for unresolved fields; validated, else Stage B is kept"),
         ("Output", "IR rendered per target LLM with token counts; Compare, History, 30-day retention")]
for k, (h, b) in enumerate(cards):
    x = 1.0 + k * 4.58
    box(s, x, 6.95, 4.25, 3.15, GREEN)
    text(s, x + 0.3, 7.15, 3.7, 0.55, h, size=26, color=MINT, font=BOLD, bold=True)
    text(s, x + 0.3, 7.8, 3.7, 2.2, b, size=18, color=MINT)
notes(s, P + "Piush. Left to right: the React UI sends prompt, target LLM, category and attachment type. Stage A detects what "
      "is missing, Stage B fixes it with rules and builds the IR. Only when the category or a reference is still unresolved "
      "(6.4% of test prompts) does Stage C, the LoRA model, run. Renderers write the same IR for Claude, GPT or Gemini; Compare "
      "runs both prompts on a real model; everything is stored through one repository module. Image mode is a separate path.")

# 4 Stage A + B ---------------------------------------------------------------------------------------------------
s = prep(3, "STAGE A DETECTS, STAGE B FIXES")
box(s, 1.0, 2.3, 7.6, 7.85, GREEN)
text(s, 1.35, 2.55, 7.0, 0.6, "STAGE A · FEATURE DETECTION", size=24, color=MINT, font=BOLD, bold=True)
bullets(s, 1.35, 3.35, 7.0, 3.6, ["Task category + confidence (embedding classifier, keyword fallback)", "A02 missing output format",
        "A03 missing length / tone / audience / language", "A04 filler and repetition", "A05 ambiguous reference (\"summarize this\")"],
        size=19, color=MINT, dot=LIGHT)
box(s, 1.35, 7.35, 6.9, 2.5, DARK)
text(s, 1.55, 7.5, 3.0, 1.2, "74.7%", size=48, color=MINT, font=ULTRA, anchor=MSO_ANCHOR.MIDDLE)
text(s, 4.55, 7.55, 3.5, 2.2, "category accuracy on 482 test prompts (coding F1 0.98, classification 0.94)", size=18, color=MINT)
text(s, 9.2, 2.35, 9.8, 0.6, "STAGE B · 15 RULES, EACH A UNIT-TESTED FUNCTION", size=24, color=GREEN, font=BOLD, bold=True)
rules = [("B01", "Remove filler"), ("B02", "Remove duplicates"), ("B03", "Output format"), ("B04", "Length limit"), ("B05", "Programming language"),
         ("B06", "Label set"), ("B07", "Structure / data block"), ("B08", "Group fallback"), ("B09-B15", "Attachment modifiers")]
for k, (c, t) in enumerate(rules):
    x, y = 9.2 + (k % 3) * 3.35, 3.15 + (k // 3) * 1.3
    box(s, x, y, 3.1, 1.1, GREEN)
    text(s, x + 0.2, y + 0.1, 2.75, 0.9, [[(c + "\n", {"font": BOLD, "bold": True, "color": MINT, "size": 19}), (t, {"size": 17})]],
         color=MINT, anchor=MSO_ANCHOR.MIDDLE, space_after=0, line=1.0)
box(s, 9.2, 7.35, 9.8, 2.8, DARK)
text(s, 9.5, 7.55, 5.4, 1.4, "3.1% → 95.0%", size=46, color=MINT, font=ULTRA, anchor=MSO_ANCHOR.MIDDLE)
text(s, 14.9, 7.6, 3.9, 2.4, "of test prompts state an output format after Stage B; 6.4% (31 of 482) are routed to Stage C", size=18, color=MINT)
notes(s, P + "Piush. Stage A on the test split: 74.7% category accuracy; coding and classification are near perfect, the "
      "closed QA / extraction / summarization group is hard on degraded prompts. Stage B: each rule is a separate unit-tested "
      "function. Example: 'write code to get all permutations of a string' gets B07 tidy, B05 'Use Python', B03 'Return only the "
      "code, in a single code block'. Result: an output format is stated in 95% of prompts instead of 3.1%. Below 0.6 "
      "confidence, B08 handles the text group without guessing the category.")

# 5 Stage C -------------------------------------------------------------------------------------------------------
s = prep(4, "STAGE C: LORA ONLY WHERE RULES STOP")
box(s, 1.0, 2.3, 8.7, 5.45, WHITE, line=PALE)
picture(s, "04_b_to_c_contract.png", 1.15, 2.45, w=8.4, border=None)
text(s, 1.0, 7.95, 8.7, 2.2, [[("B → C contract: ", {"font": BOLD, "bold": True, "color": GREEN}),
     ("Stage C gets only the unresolved fields, must return JSON with exactly those keys, is checked with Stage A's own "
      "detectors, and falls back to Stage B if invalid.", {})]], size=19)
stat(s, 10.2, 2.3, 4.3, 2.55, "5% → 98%", "answers pass validation: base model vs LoRA", size=40, lsize=18)
stat(s, 14.7, 2.3, 4.3, 2.55, "22.6 → 90.3%", "format stated on routed prompts", size=36, lsize=18)
stat(s, 10.2, 5.05, 4.3, 2.55, "0.74 s", "median per call on RTX 3050 (CPU 4.55 s)", size=44, lsize=18)
stat(s, 14.7, 5.05, 4.3, 2.55, "6.4%", "of test prompts routed to Stage C", size=44, lsize=18)
box(s, 10.2, 7.8, 8.8, 2.35, DARK)
text(s, 10.5, 7.98, 8.2, 2.0, "Forced on every prompt, Stage C lowers task intent (0.886 → 0.769), so it stays a fallback. Its "
     "category is only a suggestion: the UI asks the user and pre-selects it.", size=18, color=MINT)
notes(s, P + "Piush. Stage C is Qwen2.5-0.5B-Instruct with a LoRA adapter, trained on our 4 GB laptop GPU in 51 minutes. It runs "
      "only when Stage B leaves the category or an ambiguous reference unresolved, must answer with JSON containing exactly the "
      "requested keys, and is checked with Stage A's detectors; otherwise Stage B's result is kept. Zero-shot only 5% of answers "
      "pass, with LoRA 98.2%. On routed prompts the format is stated in 90.3% instead of 22.6%. Forcing every prompt through "
      "Stage C hurts task intent, so it stays routed-only. Median 0.74 s on GPU.")

# 6 IR / renderers / attachments / image ---------------------------------------------------------------------------
s = prep(5, "IR, RENDERERS, ATTACHMENTS, IMAGE MODE")
code = [("Claude", "<task>Write code to get\nall permutations of\na string.</task>\n<constraints>\n- Use Python.\n</constraints>\n<output_format>\nReturn only the code\n</output_format>"),
        ("GPT", "### Task\nWrite code to get\nall permutations of\na string.\n### Constraints\n- Use Python.\n### Output format\nReturn only the code"),
        ("Gemini", "Task: Write code to\nget all permutations\nof a string.\nConstraints:\nUse Python.\nOutput format:\nReturn only the code")]
for k, (h, c) in enumerate(code):
    x = 1.0 + k * 4.1
    box(s, x, 2.3, 3.85, 4.75, WHITE, line=GREEN)
    text(s, x + 0.25, 2.45, 3.4, 0.55, h, size=24, color=GREEN, font=BOLD, bold=True)
    text(s, x + 0.25, 3.05, 3.45, 3.9, c, size=14, color=DARK, font=MONO, space_after=0, line=1.0)
text(s, 1.0, 7.2, 12.1, 0.5, "Same content, layout per model · every rendering is parsed back in the tests · GPT tokens exact (tiktoken)",
     size=17, color=GREEN)
box(s, 1.0, 7.85, 12.1, 1.45, GREEN)
text(s, 1.3, 7.95, 11.5, 1.25, [[("Attachments  ", {"font": BOLD, "bold": True, "color": LIGHT}),
     ("B09-B15: image, PDF, PPTX, DOCX, spreadsheet, code, other; fire only with an attachment · ", {}),
     ("30/30", {"font": BOLD, "bold": True}), (" hand-made prompts correct", {})]], size=19, color=MINT, anchor=MSO_ANCHOR.MIDDLE)
box(s, 13.5, 2.3, 5.5, 7.0, DARK)
text(s, 13.8, 2.5, 5.0, 1.0, "IMAGE MODE (SEPARATE, USER-SELECTED)", size=21, color=MINT, font=BOLD, bold=True)
text(s, 13.8, 3.55, 5.0, 0.8, "Held-out CLIP vs the user's own prompt (30 prompts)", size=17, color=LIGHT)
for k, (a, b) in enumerate([("Original", "33.08"), ("v1: defaults", "30.78 · worse, p = 0.043"), ("v2: suggestions", "32.77 · no measurable harm")]):
    text(s, 13.8, 4.45 + k * 1.2, 5.0, 1.1, [[(a + "\n", {"font": BOLD, "bold": True}), (b, {})]], size=18, color=MINT, space_after=0)
text(s, 13.8, 8.15, 5.0, 0.9, "Style kept in 12/12 styled prompts with v2", size=17, color=LIGHT)
notes(s, P + "Nirzara. One intermediate representation, three layouts: Claude gets XML tags, GPT gets ### sections, Gemini gets "
      "plain labels with the instruction first. Tests parse every rendering back to prove no field is lost. GPT tokens are exact "
      "with tiktoken, Claude and Gemini approximate and labelled. Attachments are modifiers: B09-B15, one rule per file type, 30 "
      "of 30 hand-made prompts correct. Image mode v2 keeps the user's words and offers missing attributes as suggestions; v1, "
      "which inserted defaults, made images worse.")

# 7 Transaction flow ----------------------------------------------------------------------------------------------
s = prep(6, "TRANSACTION FLOW")
box(s, 2.6, 2.2, 14.8, 8.3, WHITE, line=PALE)
picture(s, "diag_transaction.png", 2.75, 2.3, w=14.5, border=None)
notes(s, P + "Nirzara. Follow the numbers. The endpoint first purges expired prompts. create_prompt scrubs e-mails and phone "
      "numbers and sets expires_at. Stage A results go to prompt_features, Stage B and C to optimization_results plus one "
      "transformations row per change, then renderings and token_usage per target. All through repository.py: its functions "
      "flush but never commit; the endpoint commits once, so a failure half-way stores nothing. Then Compare, History and "
      "retention close the loop.")

# 8 Database ------------------------------------------------------------------------------------------------------
s = prep(7, "DATABASE LAYER: 10 TABLES")
box(s, 1.0, 2.3, 7.2, 8.1, WHITE, line=PALE)
picture(s, "08_database_er.png", 1.15, 2.4, h=7.9, border=None)
text(s, 8.8, 2.3, 10.2, 0.6, "WRITTEN IN ONE OPTIMIZE TRANSACTION (VIA REPOSITORY.PY)", size=22, color=GREEN, font=BOLD, bold=True)
table(s, 8.8, 3.05, 10.2, [3.3, 3.2, 3.7], [["Table", "Repository call", "Stores"],
      ["prompts", "create_prompt", "PII-scrubbed text, expires_at"], ["prompt_features", "save_features", "category, confidence, gaps"],
      ["optimization_results", "save_optimization", "IR, optimized text, used_lora"], ["transformations", "save_optimization", "one row per change, rule_id"],
      ["renderings", "save_renderings", "prompt per target LLM"], ["token_usage", "save_token_usage", "tokens per target"]], size=16, row_h=0.6)
box(s, 8.8, 7.6, 10.2, 2.8, DARK)
bullets(s, 9.1, 7.8, 9.7, 2.5, ["Cascade on delete: purge removes everything derived from a prompt", "PII (e-mail, phone) scrubbed before storing",
        "30-day retention enforced before every request · SQLite / PostgreSQL"], size=18, color=MINT, dot=LIGHT, space_after=6)
notes(s, P + "Nirzara. SQLAlchemy 2.0, SQLite in development, PostgreSQL in deployment. The six tables in the table are written "
      "in one optimize transaction. Children cascade on delete, so the retention purge removes everything derived from an "
      "expired prompt. evaluation_runs is deliberately not linked to prompts so benchmark results are never purged. "
      "Constraints live in the schema: confidence between 0 and 1, a Stage B step must name a rule, one rendering per target.")

# 9 Unit testing (dark) -------------------------------------------------------------------------------------------
s = prep(8, "UNIT TESTING: 798 TESTS, ALL PASSING")
box(s, 1.0, 2.2, 10.4, 8.1, MINT)
mods = [("Image mode", 215), ("Stage B rules", 123), ("Dataset tooling", 107), ("Rendering + pipeline", 77), ("Evaluation harness", 74),
        ("Stage A", 47), ("API + UI API", 37), ("Stage C", 30), ("Database", 22), ("Coding sandbox", 21), ("Correctness suite", 17),
        ("Compare + clients", 16), ("Spelling + demo", 12)][::-1]
cd = CategoryChartData(); cd.categories = [m for m, _ in mods]; cd.add_series("tests", [v for _, v in mods])
ch = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(1.2), Inches(2.35), Inches(10.0), Inches(7.8), cd).chart
style_chart(ch, [GREEN], size=15, legend=False); ch.plots[0].gap_width = 45
ch.has_title = True; ch.chart_title.text_frame.text = "pytest tests per module (798)"
ch.chart_title.text_frame.paragraphs[0].runs[0].font.size = Pt(18); ch.chart_title.text_frame.paragraphs[0].runs[0].font.color.rgb = rgb(GREEN)
box(s, 11.9, 2.2, 7.1, 3.0, MINT)
text(s, 12.2, 2.4, 6.6, 0.55, "RUN", size=22, color=GREEN, font=BOLD, bold=True)
text(s, 12.2, 3.05, 6.6, 2.0, "cd backend\nsource .venv/bin/activate\npython -m pytest -q\n→ 798 passed in 24.6 s", size=17, font=MONO, space_after=0)
box(s, 11.9, 5.5, 7.1, 4.8, MINT)
text(s, 12.2, 5.7, 6.6, 0.55, "REAL EXAMPLES", size=22, color=GREEN, font=BOLD, bold=True)
bullets(s, 12.2, 6.35, 6.6, 3.8, ["B05 defaults to Python", "PII scrubbed before storage", "retention purge cascades",
        "blank prompt → HTTP 422, not 500", "invalid Stage C answer → Stage B kept", "sandbox has no network"], size=18, space_after=4)
notes(s, P + "Siddhi. We ran the full suite on 14 October: 798 passed in 25 seconds, offline, because providers are faked or "
      "replayed from recordings. Biggest groups: image mode 215, Stage B 123, dataset tooling 107, rendering and pipeline 77. "
      "Every Stage B rule has its own tests. Example: test_pii_is_scrubbed_before_storage stores 'mail rahul.k@example.com or "
      "call 98765 43210' and checks one [EMAIL] and two [PHONE] tags. Command: python -m pytest -q in backend.")

# 10 Integration + system ------------------------------------------------------------------------------------------
s = prep(9, "INTEGRATION AND SYSTEM TESTING")
items = [("Freeze check", "BYTE-IDENTICAL", "482 test prompts, Stage A+B vs tag frozen-for-test (sha256 35f7d4dc…)"),
         ("UI end-to-end", "13 / 13", "Playwright + Chrome: optimize, compare, suite, history, image mode, no console errors"),
         ("API end-to-end", "pytest", "FastAPI TestClient: Stage C on/off, 422 errors, retention, PII-scrubbed Compare"),
         ("Full pipeline", "482 + 44", "test split run once (2 Oct) · benchmark with a real LLM + blind judge"),
         ("Coding sandbox", "20.7 → 34.5%", "pass@1 strict, 29 Python items; bubblewrap: no network, read-only"),
         ("Correctness suite", "37 → 42 / 50", "hand-written cases with verifiable answers, gold human-reviewed")]
for k, (h, big, b) in enumerate(items):
    x, y = 1.0 + (k % 3) * 6.13, 2.3 + (k // 3) * 3.55
    box(s, x, y, 5.75, 3.3, GREEN)
    text(s, x + 0.3, y + 0.2, 5.2, 0.5, h.upper(), size=19, color=MINT, font=BOLD, bold=True)
    text(s, x + 0.3, y + 0.75, 5.2, 0.9, big, size=34, color=MINT, font=ULTRA, anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + 0.3, y + 1.75, 5.2, 1.45, b, size=18, color=MINT)
notes(s, P + "Siddhi. Beyond unit tests: the freeze check recomputes Stage A and B on all 482 test prompts and compares a sha256 "
      "with the frozen-for-test tag; on 14 October it was byte-identical. Thirteen Playwright tests drive the real app in Chrome, "
      "13 of 13 pass. API tests go end to end through FastAPI. The full pipeline ran once on the 482-prompt test split and the "
      "44-prompt benchmark. Coding answers run in a bubblewrap sandbox with no network. And the 50-case correctness suite checks real answers.")

# 11 Test cases + correctness ---------------------------------------------------------------------------------------
s = prep(10, "TEST CASES AND CORRECTNESS")
table(s, 1.0, 2.3, 11.9, [1.0, 2.7, 4.2, 2.8, 1.2], [["ID", "Case", "Input", "Result", "Status"],
      ["TC04", "B01 removes filler", "\"hey can you please summarize the text for me?\"", "\"Summarize the text.\"", "PASS"],
      ["TC16", "PII scrubbed", "mail rahul.k@example.com … 98765 43210", "1 [EMAIL], 2 [PHONE]", "PASS"],
      ["TC26", "Freeze check", "482 test prompts", "byte-identical", "PASS"],
      ["TC31", "cqa-01 hostel fee", "\"if i pay on 18 july how much total??\"", "80300 (vague said 85300)", "PASS"],
      ["TC34", "cls-09 attendance", "\"mark each student safe warning or detained\"", "got a Python script", "FAIL"],
      ["TC37", "cod-10 recipe scaler", "\"write the scale fn for recipe app\"", "4/8 asserts", "FAIL"]], size=15, row_h=0.84, status_col=4)
text(s, 1.0, 8.35, 11.9, 0.5, "Fixed 9 · broke 4 · McNemar p = 0.267 (not significant at n = 50)", size=20, color=GREEN, font=BOLD, bold=True)
text(s, 1.0, 8.95, 11.9, 1.3, "Failures: Stage A misread a table as code (cls-09); two summaries 2-3 words over the limit; one "
     "Stage C coding case passed 4 of 8 asserts.", size=18)
box(s, 13.3, 2.3, 5.7, 5.15, WHITE, line=PALE)
cd = CategoryChartData(); cd.categories = ["QA", "Extract", "Classify", "Summary", "Coding"]
cd.add_series("Vague", (8, 7, 7, 7, 8)); cd.add_series("Optimized", (10, 10, 8, 5, 9))
ch = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(13.4), Inches(2.4), Inches(5.5), Inches(4.95), cd).chart
style_chart(ch, [MAROON, GREEN], size=13); ch.value_axis.maximum_scale = 11.5
ch.has_title = True; ch.chart_title.text_frame.text = "Correct of 10, per category"
ch.chart_title.text_frame.paragraphs[0].runs[0].font.size = Pt(15); ch.chart_title.text_frame.paragraphs[0].runs[0].font.color.rgb = rgb(GREEN)
stat(s, 13.3, 7.75, 2.75, 2.45, "42/50", "optimized (vague 37)", size=36, lsize=16)
stat(s, 16.25, 7.75, 2.75, 2.45, "-29.6%", "total tokens", size=36, lsize=16)
notes(s, P + "Siddhi. The report has 38 test cases; here are six. The correctness suite: 50 new hand-written cases. Vague "
      "prompts 37 correct, optimized 42, with 29.6% fewer tokens. Optimization fixed 9 answers and broke 4; McNemar p = 0.267, "
      "so not significant at n = 50, and on Cerebras it was 43 vs 41. Be honest about the failures: cls-09, Stage A read an "
      "attendance table as code so the model wrote a script; two summaries went 2-3 words over the limit; one Stage C coding "
      "case failed 4 of 8 asserts.")

# 12 Results (dark) -----------------------------------------------------------------------------------------------
s = prep(11, "RESULTS: -39.2% TOTAL TOKENS")
box(s, 1.0, 2.2, 11.6, 8.1, MINT)
cats = ["closed_qa", "extraction", "classification", "summarization", "coding", "all"]
cd = CategoryChartData(); cd.categories = cats
cd.add_series("Vague prompt", (1115, 875, 488, 1169, 903, 909)); cd.add_series("Optimized", (417, 466, 278, 552, 324, 406))
ch = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1.2), Inches(2.35), Inches(11.2), Inches(7.8), cd).chart
style_chart(ch, [MAROON, GREEN], size=15); ch.value_axis.maximum_scale = 1300
ch.has_title = True; ch.chart_title.text_frame.text = "Mean total tokens per prompt (input + output), 482 test prompts"
ch.chart_title.text_frame.paragraphs[0].runs[0].font.size = Pt(17); ch.chart_title.text_frame.paragraphs[0].runs[0].font.color.rgb = rgb(GREEN)
stat(s, 13.1, 2.2, 5.9, 2.5, "682 → 165", "output tokens per prompt (input 226 → 241)", fill=MINT, color=DARK, size=44, lsize=18)
stat(s, 13.1, 4.95, 5.9, 2.5, "8.2 → 9.0", "benchmark quality; task success 67% → 85%", fill=MINT, color=DARK, size=44, lsize=18)
box(s, 13.1, 7.7, 5.9, 2.6, DARK)
text(s, 13.4, 7.9, 5.3, 2.2, [[("Limitation  ", {"font": BOLD, "bold": True, "color": LIGHT}),
     ("87 of 482 prompts cost more: their answer was already short (mostly extraction). Next: lean mode.", {})]], size=18, color=MINT)
notes(s, P + "Siddhi. The headline: across the 482-prompt test split, total tokens drop 39.2% per prompt, 95% CI 35.9 to 42.5%, "
      "Wilcoxon p below 0.001. The optimized prompt is a bit longer, 226 to 241 input tokens, but answers shrink from 682 to 165 "
      "because the format and length are stated. Every category saves on average; 87 of 482 short-answer prompts cost more, "
      "which is why lean mode is our first next step. On the benchmark with a blind judge, quality went 8.2 to 9.0 and task "
      "success 67 to 85%.")

# 13 Working prototype --------------------------------------------------------------------------------------------
s = prep(12, "WORKING PROTOTYPE")
picture(s, "ui_optimize.png", 1.0, 2.25, w=8.75, crop_h=7.0)
picture(s, "ui_compare-result.png", 10.25, 2.25, w=8.75, crop_h=7.0)
text(s, 1.0, 9.4, 8.75, 0.9, [[("Optimize: ", {"font": BOLD, "bold": True, "color": GREEN}),
     ("every stage shown, rule by rule, with before/after text and the rendered prompt", {})]], size=18)
text(s, 10.25, 9.4, 8.75, 0.9, [[("Compare: ", {"font": BOLD, "bold": True, "color": GREEN}),
     ("768 → 248 tokens (-67.7%) on gpt-oss-120b, both answers pass 6/6 tests", {})]], size=18)
notes(s, P + "Nirzara. FastAPI back end, React + TypeScript front end: Optimize, Compare, Test suite, Results, How it works, "
      "History, Image mode, Status. Left: Optimize on the permutations example, every stage visible with before and after text. "
      "Right: a live Compare on gpt-oss-120b: 768 tokens for the vague prompt, 248 for the optimized one, minus 67.7%, both "
      "passing 6 of 6 sandbox tests. Every result says which model actually answered; GPT and Claude need paid keys, so "
      "gpt-oss is labelled as a stand-in. The app runs on the laptop as a service and is reachable on the LAN.")

# 14 Git / GitHub -------------------------------------------------------------------------------------------------
s = prep(13, "GIT AND GITHUB")
picture(s, "gh_repo.png", 1.0, 2.35, w=9.0)
text(s, 1.0, 7.6, 9.0, 0.5, "github.com/piush365/PromptOpt", size=20, color=GREEN, font=BOLD, bold=True)
text(s, 1.0, 8.2, 9.0, 1.9, "Tags are our evaluation checkpoints: after frozen-for-test no Stage A/B rule, threshold or judge "
     "changed, and the freeze check proves it.", size=18)
stat(s, 10.6, 2.35, 2.65, 2.4, "92", "commits", size=48, lsize=18)
stat(s, 13.45, 2.35, 2.65, 2.4, "3", "branches", size=48, lsize=18)
stat(s, 16.3, 2.35, 2.7, 2.4, "4", "tags · 1 release", size=48, lsize=18)
box(s, 10.6, 5.0, 8.4, 5.1, DARK)
bullets(s, 10.9, 5.2, 7.9, 4.8, ["frozen-for-test (27 Sep): code for the final numbers", "final-for-test (2 Oct): test split run once",
        "v1.0 (8 Oct): release with the Stage C adapter", "Commits per month: Aug 7 · Sep 38 · Oct 47",
        "Tracking: roadmap and phase tables in CLAUDE.md, tags, release (no GitHub Issues/Projects used)"], size=20, color=MINT, dot=LIGHT, space_after=9)
notes(s, P + "Nirzara. All numbers are from git. 92 commits between 13 August and 10 October, on three branches: main, ui-v2 "
      "and correctness-suite. Tags were our evaluation checkpoints: frozen-for-test marks the code for the final numbers, "
      "final-for-test the one-time test run, v1.0 the release with the Stage C adapter attached. We did not use GitHub Issues or "
      "Projects; we tracked work with the roadmap and phase tables in the repository and with tags.")

# 15 Challenges (dark) --------------------------------------------------------------------------------------------
s = prep(14, "CHALLENGES AND HOW WE SOLVED THEM")
ch_ = [("Dataset generation bugs", "LLM targets dropped data → deterministic repair: 349 prompts fixed, every change logged"),
       ("Kappa paradox", "90% agreement but kappa 0.065 → report Gwet's AC1 (0.92-0.96) next to it"),
       ("API quota limits", "Groq 200k tokens/day → rolling 24-h usage ledger, 70% budget rule, resumable runs"),
       ("Stage C hallucination", "base model 5% valid → LoRA 98.2%, strict JSON contract, fallback to Stage B"),
       ("Image mode v1 failure", "defaults turned a watercolor into a photo → v2: suggestions only"),
       ("Label / regex edge cases", "\"thank you so much\" inside a button → a unit test per edge case, B16 in the app"),
       ("4 GB GPU training", "micro-batch 1 × 16, bf16, gradient checkpointing: 51 min; models loaded once")]
for k, (h, b) in enumerate(ch_):
    x, y = 1.0 + (k % 4) * 4.57, 2.3 + (k // 4) * 4.0
    box(s, x, y, 4.3, 3.75, MINT)
    text(s, x + 0.25, y + 0.2, 3.8, 0.95, h.upper(), size=19, color=GREEN, font=BOLD, bold=True, line=1.0)
    text(s, x + 0.25, y + 1.2, 3.8, 2.45, b, size=17, color=DARK)
box(s, 14.71, 6.3, 4.3, 3.75, DARK)
text(s, 14.96, 6.5, 3.8, 0.95, "REPORT IT, DON'T HIDE IT", size=19, color=MINT, font=BOLD, bold=True, line=1.0)
text(s, 14.96, 7.5, 3.8, 2.4, "After the freeze, bugs were reported, not fixed; failures are shown in the results.", size=18, color=MINT)
notes(s, P + "Nirzara. Seven problems, seven fixes. LLM-written targets dropped data, so we built a deterministic repair with a "
      "log. Raters agreed on 90% of rows but kappa was near zero, the kappa paradox, so we also report Gwet's AC1. Free API "
      "quotas: a rolling usage ledger and a 70% budget rule. Stage C invented fields zero-shot, so we fine-tuned and validate "
      "every answer. Image v1 made pictures worse; v2 suggests instead of inserting. Regex edge cases got their own tests. And "
      "the model trained on a 4 GB GPU with gradient accumulation.")

# 16 Plans --------------------------------------------------------------------------------------------------------
s = prep(15, "PLANS FOR THE UPCOMING PHASES")
plans = [("Lean mode", "add only the format line when the answer is short"), ("Smarter routing", "send label-set and typo prompts to Stage C"),
         ("Spelling normalisation", "user-confirmed fixes before optimizing"), ("Blind user tests", "prompts written by people outside the team"),
         ("Real GPT / Claude / Gemini", "Compare with paid API keys"), ("Larger correctness suite", "more cases and providers for significance"),
         ("Deployment", "PostgreSQL, HTTPS and login, beyond the LAN")]
for k, (h, b) in enumerate(plans):
    x, y = 1.0 + (k % 2) * 9.15, 2.3 + (k // 2) * 1.78
    box(s, x, y, 8.85, 1.55, GREEN)
    circle_num(s, x + 0.25, y + 0.3, 0.95, k + 1, fill=LIGHT, color=DARK, size=24)
    text(s, x + 1.45, y + 0.17, 7.2, 0.55, h, size=22, color=MINT, font=BOLD, bold=True)
    text(s, x + 1.45, y + 0.8, 7.2, 0.6, b, size=18, color=MINT)
notes(s, P + "Siddhi. In priority order. Lean mode, because 87 short-answer prompts cost more. Route label-set and typo prompts "
      "to Stage C, and let users confirm spelling fixes. Blind tests with prompts written by people outside the team. Compare on "
      "real GPT, Claude and Gemini once we have paid keys. A larger correctness suite so the gain can reach significance. Then "
      "deployment with PostgreSQL, HTTPS and login.")

# 17 Live demo (dark) ---------------------------------------------------------------------------------------------
s = prep(16, "LIVE DEMO")
box(s, 2.6, 2.3, 14.8, 7.9, MINT)
steps = [("Optimize", "\"write code to get all permutations of a string\""), ("Stage C", "\"hey can you just summarize this article for me\""),
         ("Compare", "gpt-oss-120b: tokens, latency, sandbox tests"), ("Image mode", "\"oil painting of a sailboat in a storm\""),
         ("History · Test suite · Results", "past prompts, 50-case suite, headline numbers")]
for k, (h, b) in enumerate(steps):
    y = 2.75 + k * 1.45
    circle_num(s, 3.2, y, 1.0, k + 1, fill=GREEN, color=MINT, size=26)
    text(s, 4.6, y - 0.02, 12.3, 1.1, [[(h + "   ", {"font": BOLD, "bold": True, "color": GREEN, "size": 24}), (b, {})]],
         size=21, color=DARK, anchor=MSO_ANCHOR.MIDDLE)
notes(s, P + "Piush drives the laptop; Nirzara narrates Optimize and Compare, Siddhi narrates the Test suite and Results. Follow "
      "DEMO_LA2.md. If the network fails, show the Past comparisons table on the Compare page; if the GPU fails, the app falls "
      "back to Stage A + B and says Stage C is unavailable.")

# 18 Thank you: LA-1's slide unchanged ---------------------------------------------------------------------------
notes(S[17], P + "Siddhi. Thank the panel and our guide, and invite questions.")

prs.save(OUT)
print("saved", OUT, len(prs.slides))
