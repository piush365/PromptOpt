"""Human validation of PromptOpt Dataset v1.1: rater sheets, agreement (Fleiss' Kappa) and the merged result.

Design: the three team members (TEAM) all rate the same OVERLAP set (OVERLAP_PER_CATEGORY per category, 90 in total),
which is used for Fleiss' Kappa. On top of that each team member rates EXTRA records alone, chosen by usefulness for
filtering (see `extra_priority`). The faculty rate FACULTY_PER_CATEGORY records per category (20 in total) taken
from the overlap set, so their answers can be compared with the team's majority vote (percent agreement and Cohen's
kappa): an independent check on the team's ratings. The faculty answers are not used to decide a record. The rest
of the dataset is not human-rated and relies on the automatic checks.

    python -m app.validation assign                                    # create / update the sheets
    python -m app.validation report                                    # progress, Fleiss' Kappa, faculty check
    python -m app.validation merge                                     # write the validated dataset

Everything is keyed on `source_id` (the row number in the original Dolly-15k / CodeAlpaca-20k download), which
stays the same when the notebook rebuilds the dataset; the dataset's `id` column is renumbered on every rebuild and
is not used. The assignment is frozen in `assignment.csv`: re-running `assign` keeps every overlap, faculty and
extra record already chosen and only tops up what is missing, so answers are never moved. If a pair was regenerated (its
degraded or optimized text changed), its old answers are cleared and the note says to re-rate it.
"""
import argparse
import csv
import hashlib
import math
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from pathlib import Path

from app.config import DATASET_DIR
from app.dataset_io import DEFAULT_CSV, load_rows

DEFAULT_SHEETS = DATASET_DIR.parent / "validation"
ASSIGNMENT_FILE = "assignment.csv"
TEAM = ["Nirzara Manade", "Siddhi Bolaikar", "Piush Gogi"]
FACULTY_SHEET = "faculty"
CATEGORIES = ["closed_qa", "information_extraction", "classification", "summarization", "coding"]
OVERLAP_PER_CATEGORY = 18     # records rated by the whole team, per category (90 in total), for Fleiss' Kappa
FACULTY_PER_CATEGORY = 4      # overlap records the faculty also rate, per category (20 in total)
EXTRA_PER_RATER = 60          # records each team member rates alone, for filtering
RECENT_PER_RATER = 20         # v1.2 rows (generated for the expansion) each team member rates alone, 4 per category
SEED = "promptopt-validation-v1"

# Automatic-check thresholds used when the dataset was built (DATASET_CARD.md). Pairs that passed by less than
# BORDERLINE_MARGIN are the ones the automatic checks are least sure about.
SIM_DEGRADED_MIN = 0.45
SIM_OPTIMIZED_MIN = 0.35
BORDERLINE_MARGIN = 0.10
NOISY_CATEGORIES = {"summarization", "information_extraction"}      # Dolly labels often wrong (plain questions)
SPLIT_RANK = {"benchmark": 0, "test": 1, "val": 2, "train": 3}

QUESTIONS = [
    ("Q1_degraded_same_task", "Does the DEGRADED prompt ask for the same task as the original instruction?"),
    ("Q2_degraded_realistic", "Is the DEGRADED prompt a realistic vaguer version a real user might type "
                              "(not broken, not nonsense, not a different question)?"),
    ("Q3_optimized_same_intent", "Does the OPTIMIZED prompt keep the same task and intent, without adding facts, "
                                 "answers or requirements the user did not ask for?"),
    ("Q4_optimized_better", "Is the OPTIMIZED prompt clearer and better specified than the degraded prompt "
                            "(output format, constraints, no filler)?"),
    ("Q5_category_correct", "Is the CATEGORY label correct?"),
]
QCOLS = [q for q, _ in QUESTIONS]
TEXT_COLS = ["source_id", "category", "split", "original_instruction", "context",
             "degraded_prompt", "optimized_prompt"]
SHEET_COLS = TEXT_COLS + QCOLS + ["notes"]
ANSWERS = {"Y", "N"}


# ---------------------------------------------------------------- assignment
def _unit(*parts: str) -> float:
    """Deterministic uniform number in [0, 1) from the given strings."""
    h = hashlib.sha256(":".join((SEED,) + parts).encode()).hexdigest()
    return int(h[:15], 16) / 16 ** 15


def _sim(row: dict[str, str], col: str) -> float | None:
    try:
        return float(row.get(col, ""))
    except ValueError:
        return None


def is_borderline(row: dict[str, str]) -> bool:
    """True if either similarity check passed by less than BORDERLINE_MARGIN."""
    for col, low in (("sim_degraded_vs_original", SIM_DEGRADED_MIN), ("sim_optimized_vs_original", SIM_OPTIMIZED_MIN)):
        x = _sim(row, col)
        if x is not None and x < low + BORDERLINE_MARGIN:
            return True
    return False


def sheet_name(rater: str) -> str:
    """'Nirzara Manade' -> 'Nirzara_Manade': the sheet's file name and the key used in assignment.csv."""
    return "_".join(rater.split())


def extra_priority(row: dict[str, str]) -> tuple:
    """Sort key for extra records, most useful first: evaluation splits (benchmark, test), then borderline
    automatic-check scores, then the categories with noisy Dolly labels, then a fixed pseudo-random order."""
    return (0 if row["split"] in ("benchmark", "test") else 1, 0 if is_borderline(row) else 1,
            0 if row["category"] in NOISY_CATEGORIES else 1, _unit("extra", row["source_id"]))


@dataclass
class Assignment:
    overlap: list[str] = field(default_factory=list)                    # source_ids rated by the whole team
    faculty: list[str] = field(default_factory=list)                    # subset of overlap, also rated by faculty
    single: dict[str, list[str]] = field(default_factory=dict)          # team member -> extra source_ids
    recent: dict[str, list[str]] = field(default_factory=dict)          # team member -> v1.2 source_ids

    def raters_of(self, source_id: str) -> list[str]:
        if source_id in self.overlap:
            return list(self.single) + ([FACULTY_SHEET] if source_id in self.faculty else [])
        return [s for s, ids in (*self.single.items(), *self.recent.items()) if source_id in ids]

    def role_of(self, source_id: str) -> str:
        if source_id in self.overlap:
            return "overlap"
        if any(source_id in ids for ids in self.single.values()):
            return "extra"
        return "extra_v1.2" if any(source_id in ids for ids in self.recent.values()) else "none"


def assign(rows: list[dict[str, str]], team: list[str], faculty_per_category: int = FACULTY_PER_CATEGORY,
           overlap_per_category: int = OVERLAP_PER_CATEGORY, extra: int = EXTRA_PER_RATER,
           frozen: Assignment | None = None, recent_per_rater: int = 0) -> Assignment:
    """Choose the overlap, faculty and extra records, keeping every record already in `frozen`.

    Overlap: overlap_per_category per category, fixed pseudo-random pick. Faculty: faculty_per_category per category
    from the overlap records, benchmark split first (those rows are used in the live A/B evaluation), then test.
    Extra: up to `extra` per team member, most useful first (`extra_priority`), dealt out so the shares stay equally
    useful. Recent: up to `recent_per_rater` rows per team member (spread evenly over the categories) from the rows
    generated for v1.2 (`generation_version` other than v1), same priority, none shared.
    """
    if len(set(team)) != len(team) or FACULTY_SHEET in team or not team:
        raise ValueError(f"team members must be distinct names, and not '{FACULTY_SHEET}'")
    ids = [r["source_id"] for r in rows]
    if len(set(ids)) != len(ids):
        raise ValueError("source_id is not unique in the dataset")
    present = set(ids)
    frozen = frozen or Assignment()

    by_cat = defaultdict(list)
    for r in rows:
        by_cat[r["category"]].append(r)
    cat_of = {r["source_id"]: r["category"] for r in rows}

    def top_up(chosen: list[str], per_cat: int, key, pool) -> list[str]:
        for cat in sorted(by_cat):
            have = sum(cat_of[s] == cat for s in chosen)
            ranked = sorted((r for r in by_cat[cat] if r["source_id"] not in chosen and pool(r)), key=key)
            chosen += [r["source_id"] for r in ranked[:max(per_cat - have, 0)]]
        return chosen

    overlap = top_up([s for s in frozen.overlap if s in present], overlap_per_category,
                     lambda r: _unit("overlap", r["source_id"]), lambda r: True)
    in_overlap = set(overlap)
    faculty = top_up([s for s in frozen.faculty if s in in_overlap], faculty_per_category,
                     lambda r: (SPLIT_RANK.get(r["split"], 9), _unit("faculty", r["source_id"])),
                     lambda r: r["source_id"] in in_overlap)

    taken = set(overlap)
    single = {s: [x for x in frozen.single.get(s, []) if x in present and x not in taken] for s in team}
    for ids_ in single.values():
        taken.update(ids_)
    candidates = iter(sorted((r for r in rows if r["source_id"] not in taken), key=extra_priority))
    while True:
        short = [s for s in team if len(single[s]) < extra]
        if not short:
            break
        s = min(short, key=lambda s: (len(single[s]), team.index(s)))
        nxt = next(candidates, None)
        if nxt is None:
            break
        single[s].append(nxt["source_id"])

    recent = {s: [x for x in frozen.recent.get(s, []) if x in present and x not in taken] for s in team}
    for ids_ in recent.values():
        taken.update(ids_)
    per_cat = recent_per_rater // len(CATEGORIES)
    for cat in CATEGORIES:
        pool = iter(sorted((r for r in by_cat.get(cat, []) if r["source_id"] not in taken
                            and r.get("generation_version", "v1") not in ("", "v1")), key=extra_priority))
        for s in team * per_cat:
            if sum(cat_of[x] == cat for x in recent[s]) >= per_cat:
                continue
            nxt = next(pool, None)
            if nxt is None:
                break
            recent[s].append(nxt["source_id"])
            taken.add(nxt["source_id"])
    return Assignment(overlap=overlap, faculty=faculty, single=single, recent=recent)


def read_assignment(path: Path, team: list[str]) -> Assignment:
    """The frozen assignment from a previous `assign` run. Extra records of raters not in `team` are dropped."""
    with open(path, encoding="utf-8", newline="") as f:
        rows = list(csv.DictReader(f))
    if rows and ("id" in rows[0] or any(r["role"] in ("single", "lab") for r in rows)):
        raise SystemExit(f"{path} is from an old assignment design (lab assistants or 6% overlap). Nobody had "
                         f"rated it yet, so re-run with --reset to start from the current design.")
    a = Assignment(single={s: [] for s in team}, recent={s: [] for s in team})
    dropped = set()
    for r in rows:
        if r["role"] == "overlap":
            a.overlap.append(r["source_id"])
            if FACULTY_SHEET in r["raters"].split(";"):
                a.faculty.append(r["source_id"])
        elif r["role"] in ("extra", "extra_v1.2"):
            target = a.single if r["role"] == "extra" else a.recent
            if r["raters"] in target:
                target[r["raters"]].append(r["source_id"])
            else:
                dropped.add(r["raters"])
    if dropped:
        print(f"Warning: {', '.join(sorted(dropped))} not in --team; their extra records are reassigned "
              f"(answers already given stay in their sheets).")
    return a


def write_assignment(path: Path, rows: list[dict[str, str]], a: Assignment) -> None:
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["source_id", "category", "split", "role", "raters"])
        for r in rows:
            sid = r["source_id"]
            w.writerow([sid, r["category"], r["split"], a.role_of(sid), ";".join(a.raters_of(sid))])


# ---------------------------------------------------------------- sheets (xlsx)
def _clean(value: str) -> str:
    """Text as it will round-trip through Excel: no control characters, no surrounding whitespace."""
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

    return ILLEGAL_CHARACTERS_RE.sub("", value or "").strip()


def _sheet_path(folder: Path, name: str) -> Path:
    return folder / f"{name}.xlsx"


def read_sheet(path: Path) -> list[dict[str, str]]:
    from openpyxl import load_workbook

    ws = load_workbook(path, read_only=True)["ratings"]
    it = ws.iter_rows(values_only=True)
    header = [str(h) for h in next(it)]
    return [{h: ("" if v is None else str(v)).strip() for h, v in zip(header, row)} for row in it if any(row)]


def write_sheet(path: Path, rater: str, rows: list[dict[str, str]]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = "ratings"
    ws.append(SHEET_COLS)
    for r in rows:
        ws.append([_clean(r.get(c, "")) for c in SHEET_COLS])
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if cell.data_type == "f":           # text starting with "=" must stay text, never become a formula
                cell.data_type = "s"
    widths = {"source_id": 16, "category": 20, "split": 10, "original_instruction": 45, "context": 45,
              "degraded_prompt": 45, "optimized_prompt": 60, "notes": 40}
    for i, col in enumerate(SHEET_COLS, start=1):
        letter = ws.cell(row=1, column=i).column_letter
        ws.column_dimensions[letter].width = widths.get(col, 12)
        ws.cell(row=1, column=i).font = Font(bold=True)
        if col in QCOLS:
            ws.cell(row=1, column=i).fill = PatternFill("solid", fgColor="FFF2CC")
    wrap = Alignment(wrap_text=True, vertical="top")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = wrap
    first_q = SHEET_COLS.index(QCOLS[0]) + 1
    last_q = SHEET_COLS.index(QCOLS[-1]) + 1
    dv = DataValidation(type="list", formula1='"Y,N"', allow_blank=True, showErrorMessage=True,
                        errorTitle="Y or N", error="Answer Y (yes) or N (no).")
    ws.add_data_validation(dv)
    dv.add(f"{ws.cell(row=2, column=first_q).coordinate}:{ws.cell(row=max(len(rows) + 1, 2), column=last_q).coordinate}")
    ws.freeze_panes = ws.cell(row=2, column=SHEET_COLS.index("degraded_prompt") + 1)
    ws.auto_filter.ref = ws.dimensions

    guide = wb.create_sheet("instructions")
    guide.column_dimensions["A"].width = 28
    guide.column_dimensions["B"].width = 110
    guide.append(["Rater", rater])
    guide.append(["How", "Answer every yellow column with Y or N (use the drop-down). Read the original "
                          "instruction and context first; they are the ground truth for the task."])
    guide.append(["Accept rule", "A pair is accepted only if all five answers are Y. Use 'notes' for anything "
                                 "unclear or to suggest a fix; do not edit the prompt columns."])
    guide.append([])
    for q, text in QUESTIONS:
        guide.append([q, text])
    for row in guide.iter_rows():
        for cell in row:
            cell.alignment = wrap
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def build_sheets(rows: list[dict[str, str]], a: Assignment, folder: Path) -> dict[str, dict[str, int]]:
    """Create or update one sheet per rater, keeping existing answers. Returns per-sheet statistics."""
    by_id = {r["source_id"]: r for r in rows}
    wanted = {FACULTY_SHEET: a.faculty, **{s: a.overlap + ids + a.recent.get(s, []) for s, ids in a.single.items()}}
    stats = {}
    for rater, ids in wanted.items():
        path = _sheet_path(folder, rater)
        old = {r["source_id"]: r for r in read_sheet(path)} if path.exists() else {}
        out, kept, reset, added = [], 0, 0, 0
        for sid in sorted(ids, key=lambda s: _unit("order", s)):          # mixed categories, stable order
            fresh = {c: _clean(by_id[sid].get(c, "")) for c in TEXT_COLS}
            prev = old.pop(sid, None)
            if prev is None:
                added += 1
            elif any(prev.get(c, "") != fresh[c] for c in ("degraded_prompt", "optimized_prompt")):
                reset += 1
                fresh["notes"] = "REGENERATED: pair changed since you rated it, please re-rate. " + prev.get("notes", "")
            else:
                kept += 1
                fresh.update({c: prev.get(c, "") for c in QCOLS + ["notes"]})
            out.append(fresh)
        orphaned = 0
        for sid, prev in old.items():               # never throw answers away; unanswered leftovers are dropped
            if not (any(prev.get(q) for q in QCOLS) or prev.get("notes")):
                continue
            orphaned += 1
            if not prev.get("notes", "").startswith("NO LONGER ASSIGNED"):
                prev["notes"] = ("NO LONGER ASSIGNED (removed from the dataset or reassigned; kept so no answers are "
                                 "lost). " + prev.get("notes", ""))
            out.append({c: prev.get(c, "") for c in SHEET_COLS})
        write_sheet(path, rater.replace("_", " "), out)
        stats[rater] = {"rows": len(out), "new": added, "kept": kept, "reset": reset, "orphaned": orphaned}
    return stats


# ---------------------------------------------------------------- agreement
def fleiss_kappa(counts: list[list[int]]) -> float:
    """Fleiss' Kappa. counts[i][j] = number of raters who put item i in category j (same total per item).

    Returns NaN when agreement by chance is 1 (every rating identical), where kappa is undefined.
    """
    if not counts:
        return math.nan
    n = sum(counts[0])
    if n < 2 or any(sum(r) != n for r in counts):
        raise ValueError("every item needs the same number (>= 2) of ratings")
    items = len(counts)
    p_j = [sum(r[j] for r in counts) / (items * n) for j in range(len(counts[0]))]
    p_i = [(sum(c * c for c in r) - n) / (n * (n - 1)) for r in counts]
    p_bar, p_e = sum(p_i) / items, sum(p * p for p in p_j)
    return math.nan if p_e == 1 else (p_bar - p_e) / (1 - p_e)


def interpret_kappa(k: float) -> str:
    """Landis & Koch (1977) bands."""
    if math.isnan(k):
        return "undefined (all answers identical)"
    for limit, label in ((0, "poor"), (0.2, "slight"), (0.4, "fair"), (0.6, "moderate"), (0.8, "substantial")):
        if k <= limit:
            return label
    return "almost perfect"


def _complete(r: dict[str, str]) -> bool:
    return all(r.get(q, "").upper() in ANSWERS for q in QCOLS)


def load_all(folder: Path) -> dict[str, list[dict[str, str]]]:
    sheets = {p.stem: read_sheet(p) for p in sorted(folder.glob("*.xlsx")) if not p.name.startswith("~$")}
    if not sheets:
        raise SystemExit(f"No rater sheets in {folder}. Run `python -m app.validation assign` first.")
    for name, rows in sheets.items():
        for r in rows:
            for q in QCOLS:
                v = r.get(q, "").upper()
                if v and v not in ANSWERS:
                    raise SystemExit(f"{name}.xlsx, {r['source_id']}, {q}: '{r[q]}' is not Y or N")
                r[q] = v
    return sheets


def cohen_kappa(pairs: list[tuple[str, str]]) -> float:
    """Cohen's Kappa for two raters. pairs = [(answer of rater 1, answer of rater 2), ...].

    Returns NaN when agreement by chance is 1 (both always gave the same single answer), where kappa is undefined.
    """
    if not pairs:
        return math.nan
    n = len(pairs)
    p_o = sum(a == b for a, b in pairs) / n
    labels = {x for pair in pairs for x in pair}
    p_e = sum((sum(a == c for a, _ in pairs) / n) * (sum(b == c for _, b in pairs) / n) for c in labels)
    return math.nan if p_e == 1 else (p_o - p_e) / (1 - p_e)


def _accept(r: dict[str, str]) -> str:
    return "Y" if all(r[q] == "Y" for q in QCOLS) else "N"


def team_sheets(sheets: dict[str, list[dict[str, str]]]) -> dict[str, list[dict[str, str]]]:
    return {name: rows for name, rows in sheets.items() if name != FACULTY_SHEET}


def agreement(sheets: dict[str, list[dict[str, str]]]) -> tuple[int, dict[str, tuple[float, float]]]:
    """Fleiss' Kappa and raw percent agreement per question, over records rated completely by the whole team."""
    ratings = defaultdict(dict)
    team = team_sheets(sheets)
    for rater, rows in team.items():
        for r in rows:
            if _complete(r):
                ratings[r["source_id"]][rater] = r
    shared = [by for by in ratings.values() if len(by) == len(team) >= 2]
    out = {}
    for q in QCOLS + ["accept"]:
        table = []
        for by in shared:
            vals = [_accept(r) if q == "accept" else r[q] for r in by.values()]
            table.append([vals.count("Y"), vals.count("N")])
        agree = sum(1 for y, n in table if y == 0 or n == 0) / len(table) if table else math.nan
        out[q] = (fleiss_kappa(table) if table else math.nan, agree)
    return len(shared), out


def faculty_agreement(sheets: dict[str, list[dict[str, str]]]) -> tuple[int, dict[str, tuple[float, float]]]:
    """Faculty vs the team's majority vote, per question: (Cohen's Kappa, percent agreement), over the faculty
    records that the faculty completed and the team resolved (see `resolve`)."""
    majority = resolve(team_sheets(sheets))
    faculty = [r for r in sheets.get(FACULTY_SHEET, []) if _complete(r) and r["source_id"] in majority]
    out = {}
    for q in QCOLS + ["accept"]:
        if q == "accept":
            pairs = [(_accept(r), "Y" if majority[r["source_id"]]["validation_accept"] == "True" else "N")
                     for r in faculty]
        else:
            pairs = [(r[q], majority[r["source_id"]][q]) for r in faculty]
        agree = sum(a == b for a, b in pairs) / len(pairs) if pairs else math.nan
        out[q] = (cohen_kappa(pairs), agree)
    return len(faculty), out


# ---------------------------------------------------------------- merge
def resolve(sheets: dict[str, list[dict[str, str]]]) -> dict[str, dict[str, str]]:
    """Final answers per source_id, from the team only: overlap records go by majority vote (at least 2 matching
    answers per question); extra records take their single rater's answers. The faculty sheet is ignored here: it is
    an independent check (`faculty_agreement`), not a vote."""
    by_id = defaultdict(list)
    for rater, rows in team_sheets(sheets).items():
        for r in rows:
            if _complete(r):
                by_id[r["source_id"]].append((rater, r))
    final = {}
    for sid, rated in by_id.items():
        if len(rated) == 1:
            chosen, who = {q: rated[0][1][q] for q in QCOLS}, rated[0][0]
        else:
            chosen, who = {}, "+".join(sorted(rater for rater, _ in rated))
            for q in QCOLS:
                (top, n), *rest = Counter(r[q] for _, r in rated).most_common() + [("", 0)]
                chosen[q] = top if n >= 2 and n > rest[0][1] else ""
            if not all(chosen.values()):
                continue                                                     # tie: needs another rating
        chosen["validation_accept"] = str(all(chosen[q] == "Y" for q in QCOLS))
        chosen["validated_by"] = who
        chosen["validator_notes"] = " | ".join(r["notes"] for _, r in rated if r.get("notes"))
        final[sid] = chosen
    return final


# ---------------------------------------------------------------- commands
def cmd_assign(args) -> None:
    rows = load_rows(args.dataset)
    frozen_path = args.sheets / ASSIGNMENT_FILE
    team = [sheet_name(t) for t in args.team]
    frozen = read_assignment(frozen_path, team) if frozen_path.exists() and not args.reset else None
    a = assign(rows, team, args.faculty_per_category, args.overlap_per_category, args.extra, frozen,
               args.new_per_rater)
    stats = build_sheets(rows, a, args.sheets)
    write_assignment(frozen_path, rows, a)
    extras = sum(map(len, a.single.values()))
    recent = sum(map(len, a.recent.values()))
    print(f"{len(rows)} records: {len(a.overlap)} overlap (whole team; {len(a.faculty)} of them also faculty), "
          f"{extras} extra and {recent} v1.2 (one team member each), "
          f"{len(rows) - len(a.overlap) - extras - recent} not human-rated")
    known = {sheet_name(x) for x in wanted_sheets(a)}
    stale = [p.name for p in args.sheets.glob("*.xlsx") if p.stem not in known and not p.name.startswith("~$")]
    if stale:
        print(f"Warning: {', '.join(sorted(stale))} not part of this assignment; `report` and `merge` read every "
              f"sheet in {args.sheets}, so move them out.")
    short = [c for c in CATEGORIES if sum(r["category"] == c and r["source_id"] in a.overlap for r in rows)
             < args.overlap_per_category]
    if short:
        print(f"Warning: fewer than {args.overlap_per_category} overlap records for {', '.join(short)}")
    for rater, s in stats.items():
        print(f"  {rater + '.xlsx':<24} {s['rows']:>5} rows  (new {s['new']}, unchanged {s['kept']}, "
              f"reset {s['reset']}, no longer assigned {s['orphaned']})")


def wanted_sheets(a: Assignment) -> list[str]:
    return [FACULTY_SHEET, *a.single]


def report_text(sheets: dict[str, list[dict[str, str]]]) -> str:
    lines = ["# Dataset validation report", "", "| sheet | rows | completed | accepted |", "|---|---|---|---|"]
    for name, rows in sheets.items():
        done = [r for r in rows if _complete(r)]
        acc = sum(all(r[q] == "Y" for q in QCOLS) for r in done)
        lines.append(f"| {name} | {len(rows)} | {len(done)} ({100 * len(done) / max(len(rows), 1):.0f}%) | {acc} |")
    n, stats = agreement(sheets)
    team_sets = [{r["source_id"] for r in rows} for rows in team_sheets(sheets).values()]
    n_overlap = len(set.intersection(*team_sets)) if len(team_sets) >= 2 else 0
    lines += ["", f"## Inter-rater agreement (Fleiss' Kappa, {n} of {n_overlap} overlap records rated by the whole "
                  f"team)", ""]
    if n:
        lines += ["| question | kappa | interpretation | raw agreement |", "|---|---|---|---|"]
        for q, (k, pa) in stats.items():
            lines.append(f"| {q} | {'n/a' if math.isnan(k) else f'{k:.3f}'} | {interpret_kappa(k)} | {100 * pa:.1f}% |")
        lines += ["", "When almost every answer is Y, kappa can be low even with high raw agreement "
                      "(the kappa paradox); report both."]
    else:
        lines.append("No overlap record has been completed by the whole team yet.")
    n_fac, fac = faculty_agreement(sheets)
    n_fac_rows = len(sheets.get(FACULTY_SHEET, []))
    lines += ["", f"## Faculty check (faculty vs the team's majority vote, {n_fac} of {n_fac_rows} faculty records "
                  f"compared)", ""]
    if n_fac:
        lines += ["| question | percent agreement | Cohen's kappa | interpretation |", "|---|---|---|---|"]
        for q, (k, pa) in fac.items():
            lines.append(f"| {q} | {100 * pa:.1f}% | {'n/a' if math.isnan(k) else f'{k:.3f}'} | {interpret_kappa(k)} |")
        lines += ["", "An independent check on the team's ratings: the faculty answers do not change any record."]
    else:
        lines.append("No faculty record is both completed by the faculty and resolved by the team yet.")
    final = resolve(sheets)
    accepted = sum(v["validation_accept"] == "True" for v in final.values())
    lines += ["", "## Result", "", f"Validated records: **{len(final)}**, accepted: **{accepted}**. Records nobody rated rely on the "
                                     f"automatic checks."]
    by_cat = Counter()
    cat_of = {r["source_id"]: r["category"] for rows in sheets.values() for r in rows}
    for sid, v in final.items():
        if v["validation_accept"] == "True":
            by_cat[cat_of.get(sid, "?")] += 1
    if by_cat:
        lines.append("Accepted per category: " + ", ".join(f"{c} {n}" for c, n in sorted(by_cat.items())))
    return "\n".join(lines) + "\n"


def cmd_report(args) -> None:
    text = report_text(load_all(args.sheets))
    print(text)
    (args.sheets / "validation_report.md").write_text(text, encoding="utf-8")


def cmd_merge(args) -> None:
    final = resolve(load_all(args.sheets))
    rows = load_rows(args.dataset)
    cols = list(rows[0]) + QCOLS + ["validation_accept", "validated_by", "validator_notes"]
    out_rows = []
    for r in rows:
        v = final.get(r["source_id"])
        r = dict(r, human_validated=str(v is not None), **(v or {}))
        if not args.accepted_only or (v and v["validation_accept"] == "True"):
            out_rows.append(r)
    with open(args.out, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=cols, extrasaction="ignore")
        w.writeheader()
        w.writerows(out_rows)
    print(f"Wrote {len(out_rows)} rows to {args.out} ({len(final)} validated, "
          f"{sum(v['validation_accept'] == 'True' for v in final.values())} accepted)")


def main(argv: list[str] | None = None) -> None:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--sheets", type=Path, default=DEFAULT_SHEETS, help="folder with the rater .xlsx files")
    ap.add_argument("--dataset", type=Path, default=DEFAULT_CSV)
    sub = ap.add_subparsers(dest="cmd", required=True)
    p = sub.add_parser("assign", help="create or update the rater sheets")
    p.add_argument("--team", nargs="+", default=TEAM, help="team members (one sheet each)")
    p.add_argument("--faculty-per-category", type=int, default=FACULTY_PER_CATEGORY,
                   help="overlap records per category the faculty also rate")
    p.add_argument("--overlap-per-category", type=int, default=OVERLAP_PER_CATEGORY,
                   help="records per category rated by the whole team (for Fleiss' Kappa)")
    p.add_argument("--extra", type=int, default=EXTRA_PER_RATER, help="records each team member rates alone")
    p.add_argument("--new-per-rater", type=int, default=RECENT_PER_RATER,
                   help="v1.2 rows (generation_version other than v1) each team member rates alone")
    p.add_argument("--reset", action="store_true",
                   help=f"ignore the frozen {ASSIGNMENT_FILE} and choose again (answers in the sheets are kept)")
    p.set_defaults(func=cmd_assign)
    sub.add_parser("report", help="progress, Fleiss' Kappa and the faculty check").set_defaults(func=cmd_report)
    p = sub.add_parser("merge", help="write the dataset with validation columns")
    p.add_argument("--out", type=Path, default=DATASET_DIR / "promptopt_dataset_v1_1_validated.csv")
    p.add_argument("--accepted-only", action="store_true")
    p.set_defaults(func=cmd_merge)
    args = ap.parse_args(argv)
    args.func(args)


if __name__ == "__main__":
    main()
