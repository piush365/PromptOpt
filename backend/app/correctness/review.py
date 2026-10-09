"""Human check of the gold answers: a workbook for the team, and the import of their ticks.

    python -m app.correctness.review                       # writes team_input/correctness_review/cases_review.xlsx
    python -m app.correctness.review --import filled.xlsx  # -> human_review.json + human_review.md

The workbook has an instructions sheet and 3 reviewer sheets (cases dealt round-robin within each category, so every
reviewer sees every category: 17 / 17 / 16 cases). Per case: id, category, scenario, material, vague prompt, gold
answer, how it is checked, "gold correct (Y/N)" (a drop-down) and a comment. A reviewer judges only the gold answer:
is it the one correct answer to the vague prompt, given the material?
The import reports each sheet separately and flags every case with an N; RESULTS.md lists the flagged cases.
"""
import argparse
import json
from datetime import datetime
from pathlib import Path

from app.config import BACKEND_DIR
from app.correctness.cases import CATEGORIES, SUITE_DIR, load_cases

XLSX = BACKEND_DIR.parent / "team_input" / "correctness_review" / "cases_review.xlsx"
SHEETS = ("Reviewer 1", "Reviewer 2", "Reviewer 3")
HEADERS = ("case id", "category", "scenario", "material", "vague prompt", "gold answer", "how it is checked",
           "gold correct (Y/N)", "comment")
WIDTHS = (9, 14, 30, 70, 30, 50, 34, 12, 30)


def gold_for_humans(c: dict) -> str:
    g, cat = c["gold"], c["category"]
    if cat == "closed_qa":
        return str(g)
    if cat == "information_extraction":
        return "\n".join(f"- {x}" for x in g)
    if cat == "classification":
        return "\n".join(f"{k}: {v}" for k, v in g.items())
    if cat == "summarization":
        return ("Must state:\n" + "\n".join(f"- {k['fact']}" for k in g["key_facts"])
                + "\nMust NOT state:\n" + "\n".join(f"- {f['fact']}" for f in g["forbidden"])
                + f"\nAt most {g['max_words']} words")
    return (f"Function {g['function']}; hidden asserts:\n" + "\n".join(g["tests"])
            + "\n\nReference solution:\n" + g["reference"])


def how_checked(c: dict) -> str:
    return {"closed_qa": "normalized match (wrong values in the text are also checked)",
            "information_extraction": "exact set; precision / recall / F1",
            "classification": "every item's label; per-item accuracy",
            "summarization": "blind judge ticks the key facts and forbidden statements; word count",
            "coding": "the answer's code runs the hidden asserts in the sandbox"}[c["category"]] + \
        (f". Note: {c['notes']}" if c.get("notes") else "")


def assign(cases: list[dict]) -> dict[str, list[dict]]:
    out = {s: [] for s in SHEETS}
    i = 0
    for cat in CATEGORIES:
        for c in (c for c in cases if c["category"] == cat):
            out[SHEETS[i % 3]].append(c)
            i += 1
    return out


def export(cases: list[dict], path: Path = XLSX) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    wb = Workbook()
    info = wb.active
    info.title = "Instructions"
    for row in (["PromptOpt correctness suite: check the gold answers"], [],
                ["Each reviewer fills in ONE sheet (Reviewer 1, 2 or 3)."],
                ["For each case, read the material and the vague prompt, then the gold answer."],
                ["Column 'gold correct (Y/N)': Y if the gold answer is the one correct answer to the prompt, given the "
                 "material; N if it is wrong, incomplete, or if another answer would also be correct."],
                ["If N, say why in 'comment'. Do not judge the prompt's wording, only the gold answer."],
                ["Summarization: check that the 'must state' facts are the key points and the 'must NOT state' "
                 "ones are really wrong or absent. Coding: check the asserts follow the spec in the material."]):
        info.append(row)
    info["A1"].font = Font(bold=True, size=13)
    info.column_dimensions["A"].width = 120
    for name, items in assign(cases).items():
        ws = wb.create_sheet(name)
        ws.append(HEADERS)
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="DDEBF7")
        for c in items:
            ws.append([c["id"], c["category"], c["scenario"], c["material"], c["vague_prompt"], gold_for_humans(c),
                       how_checked(c), "", ""])
        dv = DataValidation(type="list", formula1='"Y,N"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"H2:H{len(items) + 1}")
        for col, w in zip("ABCDEFGHI", WIDTHS):
            ws.column_dimensions[col].width = w
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "B2"
    wb.save(path)
    return path


def import_review(path: Path, cases: list[dict]) -> dict:
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True)
    ids = {c["id"] for c in cases}
    per_sheet, flagged, seen = {}, [], {}
    for name in SHEETS:
        if name not in wb.sheetnames:
            continue
        counts = {"Y": 0, "N": 0, "blank": 0}
        for row in wb[name].iter_rows(min_row=2, values_only=True):
            if not row or row[0] not in ids:
                continue
            tick = str(row[7] or "").strip().upper()[:1]
            tick = tick if tick in ("Y", "N") else "blank"
            counts[tick] += 1
            seen[row[0]] = tick
            if tick == "N":
                flagged.append({"id": row[0], "sheet": name, "comment": str(row[8] or "").strip()})
        per_sheet[name] = counts
    total = len(seen)
    return {"file": path.name, "imported": datetime.now().isoformat(timespec="minutes"), "per_sheet": per_sheet,
            "total": total, "yes": sum(v == "Y" for v in seen.values()), "no": sum(v == "N" for v in seen.values()),
            "blank": sum(v == "blank" for v in seen.values()), "missing": sorted(ids - set(seen)), "flagged": flagged}


def review_markdown(h: dict) -> str:
    lines = ["# Correctness suite: human review of the gold answers\n",
             f"Imported from `{h['file']}` on {h['imported']}.\n", "| sheet | Y | N | not answered |", "|---|---|---|---|"]
    lines += [f"| {s} | {c['Y']} | {c['N']} | {c['blank']} |" for s, c in h["per_sheet"].items()]
    lines += ["", f"Total: {h['yes']} Y, {h['no']} N, {h['blank']} not answered, of {h['total']} cases"
              + (f"; missing from the workbook: {', '.join(h['missing'])}" if h["missing"] else "") + ".", ""]
    lines += ["## Flagged cases (N)\n"] + ([f"- **{f['id']}** ({f['sheet']}): {f['comment'] or 'no comment'}"
                                         for f in h["flagged"]] or ["None."])
    return "\n".join(lines) + "\n"


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--import", dest="filled", type=Path, help="a filled-in cases_review.xlsx")
    args = ap.parse_args()
    cases = load_cases()
    if args.filled:
        h = import_review(args.filled, cases)
        (SUITE_DIR / "human_review.json").write_text(json.dumps(h, indent=1), encoding="utf-8")
        (SUITE_DIR / "human_review.md").write_text(review_markdown(h), encoding="utf-8")
        print(f"{h['yes']} Y, {h['no']} N, {h['blank']} blank; flagged: {[f['id'] for f in h['flagged']] or 'none'}")
    else:
        print(f"wrote {export(cases)}")


if __name__ == "__main__":
    main()
