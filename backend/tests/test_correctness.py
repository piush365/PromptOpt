"""Correctness suite: the cases, the checks, the leak check, the review workbook and the Test suite endpoints."""
import json
import shutil

import pytest
from fastapi.testclient import TestClient

from app.correctness import checks
from app.correctness.cases import CATEGORIES, by_id, check_derivable, check_suite, load_cases
from app.correctness.prompts import PROMPTS, leaks

needs_bwrap = pytest.mark.skipif(shutil.which("bwrap") is None, reason="bubblewrap not installed")
CASES = load_cases()
C = by_id(CASES)


def fixed(obj):
    return lambda system, user: obj


# ---------------------------------------------------------------- the cases
def test_fifty_cases_ten_per_category_all_valid():
    assert len(CASES) == 50
    assert {cat: sum(c["category"] == cat for c in CASES) for cat in CATEGORIES} == {cat: 10 for cat in CATEGORIES}
    non_coding = [c for c in CASES if c["category"] != "coding"]
    res = check_suite(CASES if shutil.which("bwrap") else non_coding, None)
    assert not [e for e in res["suite"] if "10 expected" not in e]
    assert {k: v for k, v in res["cases"].items() if v} == {}


@needs_bwrap
def test_every_coding_reference_passes_its_own_asserts():
    for c in CASES:
        if c["category"] == "coding":
            assert check_derivable(c) == [], c["id"]


def test_validation_catches_a_wrong_gold():
    bad = json.loads(json.dumps(C["cqa-07"]))
    bad["gold"] = "22"
    bad["check"]["aliases"] = ["22"]
    assert any("compute gives" in e for e in check_derivable(bad))
    bad = json.loads(json.dumps(C["ie-04"]))
    bad["check"]["distractors"]["Zoya"] = ["zoya"]
    assert any("not in material" in e for e in check_derivable(bad))


def test_prompts_exist_for_every_case_and_nothing_leaks():
    data = json.loads(PROMPTS.read_text())["cases"]
    assert set(data) == set(C)
    assert {k: v["leaks"] for k, v in data.items() if v["leaks"]} == {}
    for k, v in data.items():
        assert C[k]["material"] in v["optimized"]["gpt"]          # the material reaches the model
        assert v["vague"].startswith(C[k]["vague_prompt"])


def test_leak_check_finds_an_injected_answer():
    assert leaks(C["cqa-07"], "How much fine do I pay? The answer is Rs 34.") == ["gold answer '34'"]
    assert leaks(C["cls-01"], "Classify.\nR2: negative") == ["item R2 with its label 'negative'"]
    assert leaks(C["cod-02"], "Fix it.\n" + C["cod-02"]["gold"]["tests"][0])
    assert leaks(C["cqa-07"], "How much fine do I pay for the OS book?") == []


# ---------------------------------------------------------------- normalization and matching
def test_numbers_do_not_match_inside_other_numbers():
    assert checks.matches(["5"], "You can miss 5 more lectures.")
    assert not checks.matches(["5"], "You need 75% and 2.5 more")
    assert checks.matches(["80300"], "Total: **₹80,300**")
    assert not checks.matches(["1149"], "Rs 1149.50")
    assert checks.matches(["8:30"], "Take the 8:30 ordinary bus") and not checks.matches(["8:30"], "the 08:30 bus")


def test_big_o_normalization():
    n = lambda s: checks.normalize_big_o(checks.normalize(s))
    assert n("O(n log n)") == n("O(N·log N)") == "o(nlogn)"
    assert n("O(n²)") == n("O(n**2)") == n("O(n^2)") == n("O(n*n)") == "o(n^2)"
    assert n("O(log₂ n)") == n("O(log n)") == "o(logn)"


# ---------------------------------------------------------------- closed_qa
def test_closed_qa_deterministic_and_extractor():
    case = C["cqa-07"]
    assert checks.score(case, "The fine is **Rs 34**.", None)["correct"]
    assert not checks.score(case, "You pay Rs 22.", None)["correct"]
    both = "First 7 days: 14; next 4 days: 20, so not 55 but 34 in total."
    assert checks.score(case, both, None)["method"] == "unscored"
    assert checks.score(case, both, fixed({"answer": "Rs 34"}))["correct"]
    assert not checks.score(case, both, fixed({"answer": "Rs 55"}))["correct"]


# ---------------------------------------------------------------- information_extraction
def test_extraction_set_scores():
    case = C["ie-04"]
    r = checks.score(case, "- Rohan Patil\n- Vedant Joshi\n- Omkar Jadhav", None)
    assert r["correct"] and r["detail"]["f1"] == 1.0
    r = checks.score(case, "Rohan Patil and Omkar Jadhav", None)
    assert not r["correct"] and r["detail"]["recall"] == pytest.approx(0.667, abs=1e-3)
    mention = "Rohan, Vedant and Omkar. (Kunal has D grades, which are passes.)"
    r = checks.score(case, mention, fixed({"items": ["Rohan Patil", "Vedant Joshi", "Omkar Jadhav"]}))
    assert r["correct"] and r["method"] == "extractor"
    r = checks.score(case, mention, fixed({"items": ["Rohan Patil", "Vedant Joshi", "Omkar Jadhav", "Kunal More"]}))
    assert not r["correct"] and r["detail"]["precision"] == 0.75


# ---------------------------------------------------------------- classification
def test_classification_reads_lines_tables_and_headings():
    case = C["cls-04"]
    gold = case["gold"]
    lines = "\n".join(f"{k}: {v}" for k, v in gold.items())
    assert checks.score(case, lines, None)["correct"]
    table = "| Dish | Label |\n|---|---|\n" + "\n".join(f"| {k} | {v.replace('non-veg', 'Non‑Veg')} |"
                                                        for k, v in gold.items())
    assert checks.score(case, table, None)["correct"]
    heads = ("**Veg**\n- Paneer tikka\n- Dal tadka\n- Veg Manchurian\n- Mushroom masala\n- Eggless chocolate cake\n"
             "**Non-veg**\n- Egg bhurji\n- Chicken biryani\n- Fish fry\n- Omelette sandwich")
    assert checks.score(case, heads, None)["correct"]
    wrong = lines.replace("D9: veg", "D9: non-veg")
    r = checks.score(case, wrong, None)
    assert not r["correct"] and r["detail"]["wrong"] == {"D9": {"gold": "veg", "got": "non-veg"}}


def test_classification_big_o_labels_and_extractor_fallback():
    case = C["cls-08"]
    ans = "\n".join(f"{k} - {v.replace('^2', '²')}" for k, v in case["gold"].items())
    assert checks.score(case, ans, None)["correct"]
    r = checks.score(case, "S1 is constant, S2 logarithmic", fixed({"labels": {k: v for k, v in case["gold"].items()}}))
    assert r["method"] == "extractor" and r["correct"]


# ---------------------------------------------------------------- summarization
def test_summarization_checklist_and_word_limit():
    case = C["sum-04"]
    ok = {"K1": True, "K2": True, "K3": True, "K4": True, "F1": False, "F2": False, "reason": "fine"}
    short = "CSE/IT 2027 batch, 6.5 CGPA, no backlogs. Test 22 Oct, interviews 23 Oct. Register by 15 Oct, 5 pm."
    assert checks.score(case, short, fixed(ok))["correct"]
    assert not checks.score(case, short + " word" * 50, fixed(ok))["correct"]
    r = checks.score(case, short, fixed({**ok, "K3": False, "F1": True}))
    assert not r["correct"] and r["detail"]["missed"] and r["detail"]["forbidden_present"]
    assert checks.word_count("- **CSE/IT** only\n- 6.5 CGPA") == 4


# ---------------------------------------------------------------- coding
@needs_bwrap
def test_coding_reference_passes_and_buggy_code_fails():
    case = C["cod-02"]
    assert checks.score(case, "```python\n" + case["gold"]["reference"] + "```", None)["correct"]
    buggy = "```python\ndef attendance_pct(attended, total):\n    return attended // total * 100\n```"
    r = checks.score(case, buggy, None)
    assert not r["correct"] and r["detail"]["passed"] < r["detail"]["total"]


def test_mcnemar_exact():
    assert checks.mcnemar_exact(0, 0) == 1.0
    assert checks.mcnemar_exact(5, 0) == pytest.approx(0.0625)
    assert checks.mcnemar_exact(6, 1) == checks.mcnemar_exact(1, 6) == pytest.approx(0.125)


# ---------------------------------------------------------------- review workbook
def test_review_workbook_round_trip(tmp_path):
    from openpyxl import load_workbook
    from app.correctness import review
    path = review.export(CASES, tmp_path / "r.xlsx")
    wb = load_workbook(path)
    sizes = [wb[s].max_row - 1 for s in review.SHEETS]
    assert sorted(sizes) == [16, 17, 17] and sum(sizes) == 50
    ws = wb["Reviewer 2"]
    for row in range(2, ws.max_row + 1):
        ws.cell(row, 8, "Y")
    ws.cell(2, 8, "N")
    ws.cell(2, 9, "two answers possible")
    wb.save(path)
    h = review.import_review(path, CASES)
    assert h["no"] == 1 and h["yes"] == sizes[1] - 1 and h["blank"] == 50 - sizes[1]
    assert h["flagged"] == [{"id": ws.cell(2, 1).value, "sheet": "Reviewer 2", "comment": "two answers possible"}]



def test_review_import_accepts_sheets_renamed_to_reviewer_names(tmp_path):
    from openpyxl import load_workbook
    from app.correctness import review
    path = review.export(CASES, tmp_path / "r.xlsx")
    wb = load_workbook(path)
    for old, name in zip(review.SHEETS, ("Siddhi", "Nirzara", "Piush")):
        ws = wb[old]
        ws.title = name
        for row in range(2, ws.max_row + 1):
            ws.cell(row, 8, "Y")
    wb.save(path)
    h = review.import_review(path, CASES)
    assert h["yes"] == 50 and h["no"] == h["blank"] == 0 and not h["missing"]
    assert set(h["per_sheet"]) == {"Siddhi", "Nirzara", "Piush"}

# ---------------------------------------------------------------- Test suite endpoints
def test_suite_endpoints():
    from app import api
    c = TestClient(api.app)
    o = c.get("/api/suite").json()
    assert len(o["cases"]) == 50 and o["categories"] == CATEGORIES
    d = c.get("/api/suite/cqa-01", params={"model": "cerebras/gpt-oss-120b"}).json()
    assert d["case"]["material"] == C["cqa-01"]["material"] and d["gold"] == "80300"
    assert C["cqa-01"]["material"] in d["prompts"]["optimized"] and d["prompts"]["rendered_for"] == "gpt"
    assert c.get("/api/suite/nope").status_code == 404
    assert c.get("/api/suite/cqa-01", params={"model": "x/y"}).status_code == 422
    assert c.post("/api/suite/cqa-01/run", json={"model": "anthropic/claude"}).status_code == 503
