"""Dataset validation tooling: assignment, sheets, Fleiss' Kappa, merge."""
import math

import pytest

from app import validation as v

CATS = ["closed_qa", "information_extraction", "classification", "summarization", "coding"]
TEAM = ["Nirzara_Manade", "Siddhi_Bolaikar", "Piush_Gogi"]
A, B, C = TEAM


def make_rows(n_per_cat: int, start: int = 0, split=None) -> list[dict[str, str]]:
    rows = []
    for c in CATS:
        for i in range(start, start + n_per_cat):
            sp = split or ("benchmark" if i < 10 else ("test" if i < 50 else "train"))
            rows.append({"id": f"{c}-{i}", "source_id": f"src-{c}-{i}", "category": c, "split": sp,
                         "original_instruction": f"orig {c} {i}", "context": "", "degraded_prompt": f"deg {c} {i}",
                         "optimized_prompt": f"opt {c} {i}", "sim_degraded_vs_original": "0.9",
                         "sim_optimized_vs_original": "0.9"})
    return rows


def answer(row: dict, value: str = "Y") -> dict:
    return {**row, **{q: value for q in v.QCOLS}}


def test_fleiss_kappa_wikipedia_example():
    # 10 items, 14 raters, 5 categories: kappa = 0.210 (Fleiss 1971 worked example, as on Wikipedia)
    table = [[0, 0, 0, 0, 14], [0, 2, 6, 4, 2], [0, 0, 3, 5, 6], [0, 3, 9, 2, 0], [2, 2, 8, 1, 1],
             [7, 7, 0, 0, 0], [3, 2, 6, 3, 0], [2, 5, 3, 2, 2], [6, 5, 2, 1, 0], [0, 2, 2, 3, 7]]
    assert round(v.fleiss_kappa(table), 3) == 0.210


def test_fleiss_kappa_edge_cases():
    assert v.fleiss_kappa([[3, 0], [0, 3]]) == 1.0
    assert math.isnan(v.fleiss_kappa([[3, 0], [3, 0]]))          # everyone always says Y: undefined
    with pytest.raises(ValueError):
        v.fleiss_kappa([[3, 0], [2, 0]])


def test_team_names_become_sheet_names():
    assert [v.sheet_name(t) for t in v.TEAM] == TEAM


def test_assignment_sizes_and_no_record_in_two_roles():
    rows = make_rows(100)
    a = v.assign(rows, TEAM)
    cat = {r["source_id"]: r for r in rows}
    assert len(a.overlap) == 90
    assert all(sum(cat[s]["category"] == c for s in a.overlap) == 18 for c in CATS)
    assert len(a.faculty) == 20 and set(a.faculty) <= set(a.overlap)
    assert all(sum(cat[s]["category"] == c for s in a.faculty) == 4 for c in CATS)
    assert all(len(ids) == 60 for ids in a.single.values())
    flat = a.overlap + [s for ids in a.single.values() for s in ids]
    assert len(flat) == len(set(flat)) == 90 + 180
    plain = next(s for s in a.overlap if s not in a.faculty)
    assert a.raters_of(plain) == TEAM and a.role_of(plain) == "overlap"
    assert a.raters_of(a.faculty[0]) == TEAM + [v.FACULTY_SHEET] and a.role_of(a.faculty[0]) == "overlap"
    unrated = next(s for s in cat if s not in set(flat))
    assert a.raters_of(unrated) == [] and a.role_of(unrated) == "none"


def test_faculty_records_prefer_evaluation_splits():
    rows = make_rows(100)
    a = v.assign(rows, TEAM)
    cat = {r["source_id"]: r for r in rows}
    rank = lambda s: v.SPLIT_RANK[cat[s]["split"]]
    for c in CATS:
        fac = [s for s in a.faculty if cat[s]["category"] == c]
        rest = [s for s in a.overlap if cat[s]["category"] == c and s not in a.faculty]
        assert max(map(rank, fac)) <= min(map(rank, rest))


def test_extras_follow_priority_order():
    rows = make_rows(100)
    a = v.assign(rows, TEAM, extra=10)
    cat = {r["source_id"]: r for r in rows}
    assert all(cat[s]["split"] in ("benchmark", "test") for ids in a.single.values() for s in ids)

    rows = make_rows(40, split="train")                       # same split: borderline first, then noisy categories
    for r in rows[::7]:
        r["sim_optimized_vs_original"] = "0.40"               # passed the 0.35 check by less than the margin
    borderline = {r["source_id"] for r in rows[::7]}
    a = v.assign(rows, TEAM, extra=4)
    extras = [s for ids in a.single.values() for s in ids]
    assert set(extras) <= borderline
    left = borderline - set(a.overlap)
    a = v.assign(rows, TEAM, extra=len(left) // 3 + 3)    # borderline ones used up, noisy categories next
    cat = {r["source_id"]: r for r in rows}
    extras = {s for ids in a.single.values() for s in ids}
    rest = extras - borderline
    assert left <= extras and rest and all(cat[s]["category"] in v.NOISY_CATEGORIES for s in rest)


def test_extras_are_dealt_evenly_by_usefulness():
    rows = make_rows(100)
    a = v.assign(rows, TEAM)
    ranked = sorted((r for r in rows if r["source_id"] not in set(a.overlap)), key=v.extra_priority)
    first = [r["source_id"] for r in ranked[:3]]
    assert [a.single[s][0] for s in TEAM] == first        # every student gets one of the top 3


def test_small_dataset_gives_what_it_can():
    a = v.assign(make_rows(10), TEAM)                         # 50 rows: 50 overlap, nothing left
    assert len(a.overlap) == 50 and len(a.faculty) == 20 and all(ids == [] for ids in a.single.values())


def test_frozen_assignment_survives_growth():
    small, big = make_rows(30), make_rows(30) + make_rows(70, start=30)
    a = v.assign(small, TEAM, extra=20)
    b = v.assign(big, TEAM, extra=30, frozen=a)
    assert b.faculty == a.faculty
    assert b.overlap[:len(a.overlap)] == a.overlap and len(b.overlap) == 90
    for s in TEAM:
        assert a.single[s] and b.single[s][:len(a.single[s])] == a.single[s] and len(b.single[s]) == 30
    unfrozen = v.assign(big, TEAM, extra=30)
    assert set(unfrozen.overlap) != set(b.overlap)            # without freezing, the pick would move


def test_frozen_rows_that_left_the_dataset_are_dropped():
    rows = make_rows(100)
    a = v.assign(rows, TEAM)
    gone = a.overlap[0]
    b = v.assign([r for r in rows if r["source_id"] != gone], TEAM, frozen=a)
    assert gone not in b.overlap and len(b.overlap) == 90 and b.overlap[:89] == a.overlap[1:]


def test_assignment_file_round_trip(tmp_path):
    rows = make_rows(100)
    a = v.assign(rows, TEAM)
    path = tmp_path / v.ASSIGNMENT_FILE
    v.write_assignment(path, rows, a)
    b = v.read_assignment(path, TEAM)
    assert set(b.faculty) == set(a.faculty) and set(b.overlap) == set(a.overlap)
    assert all(set(b.single[s]) == set(a.single[s]) for s in TEAM)
    assert v.read_assignment(path, [A, B, "new"]).single["new"] == []


@pytest.mark.parametrize("content", [
    "source_id,id,category,split,role,raters\nsrc-1,PO-1,coding,train,single,student_A\n",
    "source_id,category,split,role,raters\nsrc-1,coding,benchmark,lab,lab_assistants\n",
])
def test_old_assignment_file_is_refused(tmp_path, content):
    path = tmp_path / v.ASSIGNMENT_FILE
    path.write_text(content)
    with pytest.raises(SystemExit, match="--reset"):
        v.read_assignment(path, TEAM)


def test_assignment_rejects_bad_input():
    with pytest.raises(ValueError):
        v.assign(make_rows(20), ["a", "a", "b"])
    rows = make_rows(20)
    rows[1]["source_id"] = rows[0]["source_id"]
    with pytest.raises(ValueError):
        v.assign(rows, TEAM)


def test_sheets_are_keyed_by_source_id_not_id():
    assert "id" not in v.SHEET_COLS and "source_id" in v.SHEET_COLS


def test_sheets_keep_answers_and_reset_regenerated_pairs(tmp_path):
    rows = make_rows(30)
    a = v.assign(rows, TEAM, extra=10)
    v.build_sheets(rows, a, tmp_path)
    sheet = tmp_path / f"{A}.xlsx"
    rated = v.read_sheet(sheet)
    rated[0] = answer(rated[0], "Y")
    rated[1] = answer(rated[1], "N") | {"notes": "drifted"}
    v.write_sheet(sheet, A, rated)
    kept_id, regen_id = rated[0]["source_id"], rated[1]["source_id"]

    grown = make_rows(30) + make_rows(20, start=30)
    for r in grown:
        r["id"] = "RENUMBERED-" + r["id"]                     # the notebook renumbers `id` on every rebuild
        if r["source_id"] == regen_id:
            r["optimized_prompt"] = "a regenerated optimized prompt"
    stats = v.build_sheets(grown, v.assign(grown, TEAM, extra=20, frozen=a), tmp_path)

    after = {r["source_id"]: r for r in v.read_sheet(sheet)}
    assert all(after[kept_id][q] == "Y" for q in v.QCOLS)
    assert all(after[regen_id][q] == "" for q in v.QCOLS)
    assert after[regen_id]["notes"].startswith("REGENERATED") and "drifted" in after[regen_id]["notes"]
    assert stats[A]["kept"] >= 1 and stats[A]["reset"] == 1 and stats[A]["new"] > 0


def test_sheets_drop_unanswered_rows_and_keep_answered_ones_no_longer_assigned(tmp_path):
    rows = make_rows(30)
    v.build_sheets(rows, v.assign(rows, TEAM, extra=10), tmp_path)
    sheet = tmp_path / f"{A}.xlsx"
    rated = v.read_sheet(sheet)
    rated[0] = answer(rated[0], "Y")
    v.write_sheet(sheet, A, rated)

    empty = v.Assignment(single={s: [] for s in TEAM})     # e.g. after --reset with a different pick
    stats = v.build_sheets(rows, empty, tmp_path)
    after = v.read_sheet(sheet)
    assert [r["source_id"] for r in after] == [rated[0]["source_id"]]
    assert after[0]["notes"].startswith("NO LONGER ASSIGNED") and after[0]["Q1_degraded_same_task"] == "Y"
    assert stats[A]["orphaned"] == 1 and stats[B]["rows"] == 0


def test_text_starting_with_equals_stays_text(tmp_path):
    row = make_rows(1)[0] | {"degraded_prompt": "=SUM(A1:A3) what does this do\x07"}
    path = tmp_path / "x.xlsx"
    v.write_sheet(path, "x", [row])
    assert v.read_sheet(path)[0]["degraded_prompt"] == "=SUM(A1:A3) what does this do"
    from openpyxl import load_workbook
    ws = load_workbook(path)["ratings"]
    col = v.SHEET_COLS.index("degraded_prompt") + 1
    assert ws.cell(row=2, column=col).data_type == "s"              # stored as text, not a formula


def _sheets_with_overlap(answers_by_student):
    """One overlap record rated by every student with the given answer for all questions."""
    base = make_rows(1)[0]
    return {s: [answer(base, a)] for s, a in answers_by_student.items()}


def test_resolve_majority_vote_ignores_faculty():
    sheets = _sheets_with_overlap({A: "Y", B: "Y", C: "N"})
    final = v.resolve(sheets)
    (res,) = final.values()
    assert res["validation_accept"] == "True" and res["validated_by"] == "+".join(sorted(TEAM))

    faculty_row = answer(make_rows(1)[0], "N")                      # same record: the faculty disagree
    assert v.resolve({**sheets, v.FACULTY_SHEET: [faculty_row]}) == final
    only_faculty = answer(make_rows(1)[0] | {"source_id": "fac-only"}, "N")
    assert "fac-only" not in v.resolve({**sheets, v.FACULTY_SHEET: [only_faculty]})


def test_two_way_tie_stays_unresolved():
    sheets = _sheets_with_overlap({A: "Y", B: "N"})
    sheets[C] = []
    assert v.resolve(sheets) == {}


def test_agreement_and_report(tmp_path):
    rows = [answer(r) for r in make_rows(4)]
    sheets = {s: [dict(r) for r in rows] for s in TEAM}
    sheets[C][0] = answer(rows[0], "N")
    n, stats = v.agreement(sheets)
    assert n == len(rows)
    kappa, raw = stats["Q1_degraded_same_task"]
    assert raw == pytest.approx((len(rows) - 1) / len(rows))
    text = v.report_text(sheets)
    assert "Fleiss" in text and f"Validated records: **{len(rows)}**" in text


def test_invalid_answer_is_reported(tmp_path):
    row = answer(make_rows(1)[0]) | {"Q1_degraded_same_task": "maybe"}
    v.write_sheet(tmp_path / f"{A}.xlsx", A, [row])
    with pytest.raises(SystemExit, match="not Y or N"):
        v.load_all(tmp_path)


def test_cohen_kappa():
    # Wikipedia worked example: 50 items, 20 YY, 5 YN, 10 NY, 15 NN -> p_o 0.7, p_e 0.5, kappa 0.4
    pairs = [("Y", "Y")] * 20 + [("Y", "N")] * 5 + [("N", "Y")] * 10 + [("N", "N")] * 15
    assert v.cohen_kappa(pairs) == pytest.approx(0.4)
    assert v.cohen_kappa([("Y", "Y"), ("N", "N")]) == 1.0
    assert math.isnan(v.cohen_kappa([("Y", "Y"), ("Y", "Y")]))       # both always Y: undefined
    assert math.isnan(v.cohen_kappa([]))


def test_faculty_agreement_against_team_majority():
    rows = [answer(r) for r in make_rows(2)]                            # 10 records, all Y from the team
    sheets = {s: [dict(r) for r in rows] for s in TEAM}
    sheets[C][0] = answer(rows[0], "N")                                 # outvoted: majority still Y
    for s in (A, B):                                                    # record 1: team majority N on Q2
        sheets[s][1] = rows[1] | {"Q2_degraded_realistic": "N"}
    faculty = [dict(r) for r in rows[:4]]
    faculty[1] = rows[1] | {"Q2_degraded_realistic": "N"}               # agrees with the team's N
    faculty[2] = rows[2] | {"Q4_optimized_better": "N"}                 # disagrees with the team's Y
    faculty.append(answer(make_rows(1)[0] | {"source_id": "not-rated-by-team"}, "Y"))
    faculty.append(rows[4] | {"Q1_degraded_same_task": ""})             # incomplete: skipped
    sheets[v.FACULTY_SHEET] = faculty

    n, stats = v.faculty_agreement(sheets)
    assert n == 4
    assert stats["Q1_degraded_same_task"][1] == 1.0
    assert stats["Q2_degraded_realistic"] == (pytest.approx(1.0), 1.0)
    k4, pa4 = stats["Q4_optimized_better"]
    assert pa4 == 0.75 and k4 == pytest.approx(0.0)                     # team always Y: no better than chance
    assert stats["accept"][1] == pytest.approx(0.75)                    # record 2: faculty reject, team accept

    text = v.report_text(sheets)
    assert "Faculty check" in text and "4 of 6 faculty records" in text and "Cohen's kappa" in text
    n, _ = v.agreement(sheets)
    assert n == len(rows)                                               # the faculty sheet is not a team rater
